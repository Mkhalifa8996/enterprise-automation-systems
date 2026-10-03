# -*- coding: utf-8 -*-
"""نسخة إكسل المرآة: قراءة ملف data.xlsx القديم واستيراده، وكتابة نسخة محدثة.

قاعدة البيانات هي مصدر الحقيقة. هذا الملف يبقي `data.xlsx` محدَّثاً ومقروءاً
في برنامج إكسل كما كان سابقاً، ويحفظ طريقة ترميز تسميات البنود المخصّصة داخل
عمود الملاحظات على شكل JSON موسوم، وهي نفس الطريقة المستخدمة في النسخة
القديمة من البرنامج (لضمان قراءة الملفات المنشأة سابقاً دون فقد).
"""

import json
import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .. import config
from ..constants import (
    CUSTOMER_HEADERS,
    INVOICE_HEADERS,
    META_FIELDS,
    PAYMENT_HEADERS,
    SERVICE_ITEMS,
    SHEET_CUSTOMERS,
    SHEET_INVOICES,
    SHEET_PAYMENTS,
)
from ..utils import safe_int, safe_str
from .db import SaveError

HEADER_FILL = PatternFill(start_color="0F2B46", end_color="0F2B46", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")

LABELS_MARKER = "__SERVICE_LABELS__"


# ------------------------------------------------------------------
# ترميز/فك ترميز تسميات البنود المخصّصة داخل عمود الملاحظات
# ------------------------------------------------------------------
def _default_service_labels_payload():
    return json.dumps([{"ar": ar, "en": en} for ar, en in SERVICE_ITEMS], ensure_ascii=False)


def _encode_service_labels(items):
    """يُنشئ نسخة JSON من تسميات الخدمات الحالية."""
    labels = []
    for idx, (ar, en) in enumerate(SERVICE_ITEMS):
        item = items[idx] if items and idx < len(items) else {}
        ar_label = item.get("label") if item else None
        en_label = item.get("labelEn") if item else None
        labels.append({
            "ar": ar_label if ar_label not in (None, "") else ar,
            "en": en_label if en_label not in (None, "") else en,
        })
    return json.dumps(labels, ensure_ascii=False)


def _encode_service_labels_for_notes(items):
    """يعيد التسمية المُرمَّزة إن اختلفت عن الافتراضية، وإلا نصاً فارغاً."""
    payload = _encode_service_labels(items)
    if payload == _default_service_labels_payload():
        return ""
    return payload


def _decode_service_labels(value):
    if not isinstance(value, str):
        return None
    if value.startswith(LABELS_MARKER):
        value = value[len(LABELS_MARKER):]
    try:
        decoded = json.loads(value)
    except (TypeError, ValueError):
        return None
    return decoded if isinstance(decoded, list) else None


def _split_service_labels_from_notes(notes_value):
    """يفصل نص الملاحظات عن تسمية البنود المخصّصة المدمجة فيه."""
    if not isinstance(notes_value, str):
        return notes_value, None
    if LABELS_MARKER not in notes_value:
        return notes_value, None
    parts = notes_value.split(LABELS_MARKER, 1)
    if len(parts) != 2:
        return notes_value, None
    notes_text = parts[0].rstrip()
    payload = parts[1]
    if payload.endswith("__"):
        payload = payload[:-2]
    return notes_text, _decode_service_labels(payload)


# ------------------------------------------------------------------
# التحويل بين صف إكسل وقاموس فاتورة
# ------------------------------------------------------------------
def row_to_invoice_dict(row_vals):
    """يحوّل صفاً من ورقة الفواتير إلى قاموس، مع فك تسميات البنود المخصّصة."""
    d = {}
    for i, (key, _label) in enumerate(META_FIELDS):
        value = row_vals[i] if i < len(row_vals) else None
        d[key] = "" if value is None else value
    base = len(META_FIELDS)
    items = []
    labels_data = None
    notes_idx = base + len(SERVICE_ITEMS) * 2
    notes_raw = row_vals[notes_idx] if len(row_vals) > notes_idx else None
    notes_value, labels_data = _split_service_labels_from_notes(notes_raw)
    d["notes"] = notes_value if notes_value is not None else ""
    for idx, (ar, en) in enumerate(SERVICE_ITEMS):
        dinar = row_vals[base + idx * 2] if base + idx * 2 < len(row_vals) else ""
        fils = row_vals[base + idx * 2 + 1] if base + idx * 2 + 1 < len(row_vals) else ""
        item_label, item_label_en = ar, en
        if labels_data and idx < len(labels_data):
            item_label = labels_data[idx].get("ar", ar)
            item_label_en = labels_data[idx].get("en", en)
        items.append({
            "label": item_label, "labelEn": item_label_en,
            "dinar": dinar if dinar is not None else "",
            "fils": fils if fils is not None else "",
        })
    d["items"] = items
    d["total"] = row_vals[notes_idx + 1] if notes_idx + 1 < len(row_vals) else 0
    return d


def invoice_dict_to_row(d):
    """يحوّل قاموس فاتورة إلى صف جاهز للكتابة في ورقة الفواتير."""
    row = [d.get(key, "") for key, _label in META_FIELDS]
    total_fils = 0
    for item in d.get("items") or []:
        dinar = safe_int(item.get("dinar"))
        fils = safe_int(item.get("fils"))
        row.append(dinar if dinar else "")
        row.append(fils if fils else "")
        total_fils += dinar * 1000 + fils
    notes_value = d.get("notes", "") or ""
    encoded_labels = _encode_service_labels_for_notes(d.get("items") or [])
    if encoded_labels:
        notes_value = (
            f"{notes_value}\n{LABELS_MARKER}{encoded_labels}__"
            if notes_value else f"{LABELS_MARKER}{encoded_labels}__"
        )
    row.append(notes_value)
    row.append(round(total_fils / 1000, 3))
    return row
# ------------------------------------------------------------------
# قراءة ملف إكسل (للاستيراد الأولي) وكتابته (كمرآة)
# ------------------------------------------------------------------
def _write_header(ws, headers):
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=i, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"


def _safe_save(wb):
    """يحفظ ملف العمل، ويحوّل أخطاء القرص الشائعة إلى رسالة عربية واضحة."""
    try:
        wb.save(config.DATA_FILE)
    except PermissionError:
        raise SaveError(
            "تعذّر تحديث ملف data.xlsx.\n"
            "الرجاء إغلاق الملف إذا كان مفتوحاً في برنامج إكسل ثم المحاولة مرة أخرى.\n"
            "(بياناتك محفوظة في قاعدة البيانات ولن تضيع.)"
        )
    except OSError as e:
        raise SaveError(f"تعذّر تحديث ملف data.xlsx:\n{e}")


def read_workbook_data():
    """يقرأ ملف data.xlsx ويعيد (فواتير، عملاء، دفعات) كقوائم قواميس.

    يعيد قوائم فارغة إن لم يكن الملف موجوداً أو كان بلا الأوراق المتوقعة،
    حتى يبقى الاستيراد الأولي آمناً على ملف ناقص أو تالف.
    """
    invoices, customers, payments = [], [], []
    if not os.path.exists(config.DATA_FILE):
        return invoices, customers, payments
    try:
        wb = load_workbook(config.DATA_FILE, data_only=True)
    except Exception:
        # ملف تالف أو غير قابل للقراءة: لا نفقد شيئاً، فقط لا نستورد منه.
        return invoices, customers, payments

    if SHEET_INVOICES in wb.sheetnames:
        ws = wb[SHEET_INVOICES]
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=1).value in (None, ""):
                continue
            values = [ws.cell(row=r, column=c).value
                      for c in range(1, len(INVOICE_HEADERS) + 1)]
            invoices.append(row_to_invoice_dict(values))

    if SHEET_CUSTOMERS in wb.sheetnames:
        ws = wb[SHEET_CUSTOMERS]
        for r in range(2, ws.max_row + 1):
            name = ws.cell(row=r, column=1).value
            if name in (None, ""):
                continue
            customers.append({
                "name": name,
                "phone": ws.cell(row=r, column=2).value or "",
                "address": ws.cell(row=r, column=3).value or "",
                "notes": ws.cell(row=r, column=4).value or "",
                "company_code": safe_str(ws.cell(row=r, column=5).value),
            })

    if SHEET_PAYMENTS in wb.sheetnames:
        ws = wb[SHEET_PAYMENTS]
        for r in range(2, ws.max_row + 1):
            pid = ws.cell(row=r, column=1).value
            if pid in (None, ""):
                continue
            applied_raw = ws.cell(row=r, column=6).value or ""
            applied = [x.strip() for x in str(applied_raw).split(",") if x.strip()]
            payments.append({
                "id": pid,
                "customer": ws.cell(row=r, column=2).value or "",
                "date": ws.cell(row=r, column=3).value or "",
                "amount": ws.cell(row=r, column=4).value or 0,
                "notes": ws.cell(row=r, column=5).value or "",
                "applied_invoices": applied,
            })
    return invoices, customers, payments


def workbook_has_data():
    """هل يحتوي ملف data.xlsx أي بيانات فعلاً؟ (يُستخدم لتحديد أول استيراد)"""
    invoices, customers, payments = read_workbook_data()
    return bool(invoices or customers or payments)


def export_to_excel(invoices, customers, payments):
    """يعيد كتابة ملف data.xlsx كاملاً من محتوى قاعدة البيانات."""
    wb = Workbook()
    ws1 = wb.active
    ws1.title = SHEET_INVOICES
    _write_header(ws1, INVOICE_HEADERS)
    for inv in invoices:
        ws1.append(invoice_dict_to_row(inv))

    ws2 = wb.create_sheet(SHEET_CUSTOMERS)
    _write_header(ws2, CUSTOMER_HEADERS)
    for c in customers:
        ws2.append([c.get("name", ""), c.get("phone", ""), c.get("address", ""),
                    c.get("notes", ""), c.get("company_code", "")])

    ws3 = wb.create_sheet(SHEET_PAYMENTS)
    _write_header(ws3, PAYMENT_HEADERS)
    for p in payments:
        applied = p.get("applied_invoices") or []
        ws3.append([p.get("id", ""), p.get("customer", ""), p.get("date", ""),
                    round(float(p.get("amount") or 0), 3), p.get("notes", ""),
                    ", ".join(str(x).strip() for x in applied if str(x).strip())])

    _safe_save(wb)