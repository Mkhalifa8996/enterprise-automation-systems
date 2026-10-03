# -*- coding: utf-8 -*-
"""اختبارات الترحيل من إكسل، والترقيم التلقائي، والحسابات، ونسخة إكسل المرآة."""

import os
import sys

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from customs_billing.constants import (
    CUSTOMER_HEADERS,
    INVOICE_HEADERS,
    PAYMENT_HEADERS,
    SERVICE_ITEMS,
    SHEET_CUSTOMERS,
    SHEET_INVOICES,
    SHEET_PAYMENTS,
)


@pytest.fixture()
def data(tmp_path, monkeypatch):
    import customs_billing.config as config
    import customs_billing.data as data_pkg

    monkeypatch.setattr(config, "DB_FILE", str(tmp_path / "test.db"))
    monkeypatch.setattr(config, "DATA_FILE", str(tmp_path / "data.xlsx"))
    monkeypatch.setattr(data_pkg, "_migration_done", False, raising=False)
    return data_pkg


def legacy_workbook(path, invoices=None, customers=None, payments=None):
    """يبني ملف data.xlsx بالشكل القديم (كما كان يولّده الإصدار السابق)."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_INVOICES
    ws.append(INVOICE_HEADERS)
    for row in invoices or []:
        ws.append(row)

    ws2 = wb.create_sheet(SHEET_CUSTOMERS)
    ws2.append(CUSTOMER_HEADERS)
    for row in customers or []:
        ws2.append(row)

    ws3 = wb.create_sheet(SHEET_PAYMENTS)
    ws3.append(PAYMENT_HEADERS)
    for row in payments or []:
        ws3.append(row)

    wb.save(path)


def legacy_invoice_row(invno="010001", date="2026-07-11", customer="East Port",
                       dinar=0, fils=0, filled=1):
    """صف فاتورة بطول الأعمدة القديم: 9 حقول + 48 عمود بند + ملاحظات + إجمالي.

    `filled` يحدد عدد البنود التي تحمل قيمة؛ الباقي فارغ تماماً كما في
    ملفات الإصدارات السابقة.
    """
    row = [invno, date, customer, "434343", "2026-07-10", "الشعب", "2", "بضائع", "الكويت"]
    for i in range(len(SERVICE_ITEMS)):
        row.extend([dinar if i < filled else "", fils if i < filled else ""])
    total = filled * (dinar * 1000 + fils) / 1000
    row.extend(["", total])
    return row


# ------------------------------------------------------------------ الترحيل
def test_migration_imports_legacy_workbook(data, tmp_path):
    legacy_workbook(
        tmp_path / "data.xlsx",
        invoices=[legacy_invoice_row("010001", dinar=10, fils=250),
                  legacy_invoice_row("010002", dinar=5)],
        customers=[["East Port", "123", "عنوان", "ملاحظة", "1"],
                   ["Majed", "456", "عنوان", "", "2"]],
        payments=[["010001", "East Port", "2026-07-30", 15.25, "دفعة", "010001"]],
    )

    assert data.migrate_excel_to_database() is True

    customers = data.load_customers()
    assert sorted(c["name"] for c in customers) == ["East Port", "Majed"]
    # رمز الشركة يُطبَّع أثناء الترحيل
    assert {c["name"]: c["company_code"] for c in customers}["East Port"] == "01"

    invoices = data.load_invoices()
    assert len(invoices) == 2
    first = next(i for i in invoices if i["invno"] == "010001")
    assert first["total"] == pytest.approx(10.25)
    assert len(first["items"]) == len(SERVICE_ITEMS)
    assert first["items"][0]["dinar"] == 10

    payments = data.load_payments()
    assert payments[0]["id"] == "010001"
    assert payments[0]["applied_invoices"] == ["010001"]


def test_migration_skipped_when_db_not_empty(data, tmp_path):
    legacy_workbook(tmp_path / "data.xlsx",
                    invoices=[legacy_invoice_row("010001", dinar=1)],
                    customers=[["East Port", "", "", "", "1"]])
    assert data.migrate_excel_to_database() is True

    # الاستيراد الثاني يجب ألا يكرّر البيانات
    assert data.migrate_excel_to_database() is False
    assert len(data.load_invoices()) == 1


def test_migration_noop_without_legacy_file(data):
    assert data.migrate_excel_to_database() is False
    assert data.load_invoices() == []
# ------------------------------------------------------------------ الترقيم
def test_company_codes_normalized(data):
    assert data.normalize_company_code("1") == "01"
    assert data.normalize_company_code(7) == "07"
    assert data.normalize_company_code("99") == "99"
    assert data.normalize_company_code("0") == ""
    assert data.normalize_company_code("100") == ""
    assert data.normalize_company_code("abc") == ""
    assert data.normalize_company_code("") == ""
    assert data.normalize_company_code(None) == ""


def test_next_available_company_code(data):
    assert data.next_available_company_code() == "01"
    data.save_customer({"name": "A", "company_code": "01"})
    assert data.next_available_company_code() == "02"


def test_next_invoice_number_sequences_per_company(data):
    assert data.next_invoice_number("01") == "010001"
    data.save_invoice({
        "invno": "010001", "date": "2026-01-01", "customer": "A",
        "items": [{"label": "", "labelEn": "", "dinar": 1, "fils": 0}] * 24,
    })
    assert data.next_invoice_number("01") == "010002"
    # شركة أخرى تبدأ تسلسلياً من 01
    assert data.next_invoice_number("02") == "020001"


def test_next_invoice_number_rejects_bad_code(data):
    with pytest.raises(ValueError):
        data.next_invoice_number("0")
    with pytest.raises(ValueError):
        data.next_invoice_number("100")


def test_next_receipt_number_sequences(data):
    assert data.next_receipt_number("01") == "010001"
    data.save_payment({"id": "010001", "customer": "A", "amount": 1})
    assert data.next_receipt_number("01") == "010002"


def test_find_customer_by_company_code(data):
    data.save_customer({"name": "East Port", "company_code": "01"})
    data.save_customer({"name": "Majed", "company_code": "02"})
    assert data.find_customer_by_company_code("01") == "East Port"
    assert data.find_customer_by_company_code("01", exclude_name="East Port") is None
    assert data.find_customer_by_company_code("77") is None


# ------------------------------------------------------------------ الحسابات
def _invoice(invno, total, customer="East Port", date="2026-01-01"):
    return {"invno": invno, "total": total, "customer": customer, "date": date,
            "items": []}


def test_customer_balances(data):
    invs = [_invoice("1", 100.0), _invoice("2", 50.0, customer="Majed")]
    pays = [{"id": "p1", "customer": "East Port", "amount": 30.0, "applied_invoices": []}]
    balances = {b["customer"]: b for b in data.customer_balances(invs, pays)}
    assert balances["East Port"]["balance"] == pytest.approx(70.0)
    assert balances["Majed"]["balance"] == pytest.approx(50.0)


def test_invoice_payment_status_specific_allocation(data):
    invs = [_invoice("1", 100.0), _invoice("2", 100.0)]
    pays = [{"id": "p1", "customer": "East Port", "amount": 60.0,
             "applied_invoices": ["2"]}]
    status = data.invoice_payment_status(invs, pays)
    # كل المبلغ يذهب إلى الفاتورة المحددة أولاً
    assert status["2"]["paid"] == pytest.approx(60.0)
    assert status["2"]["status"] == "partial"
    assert status["1"]["paid"] == pytest.approx(0.0)
    assert status["1"]["status"] == "unpaid"


def test_invoice_payment_status_fifo_for_general_payments(data):
    invs = [_invoice("1", 100.0, date="2026-01-01"),
            _invoice("2", 100.0, date="2026-02-01")]
    pays = [{"id": "p1", "customer": "East Port", "amount": 150.0, "applied_invoices": []}]
    status = data.invoice_payment_status(invs, pays)
    # الأقدم أولاً
    assert status["1"]["paid"] == pytest.approx(100.0)
    assert status["1"]["status"] == "paid"
    assert status["2"]["paid"] == pytest.approx(50.0)
    assert status["2"]["status"] == "partial"


def test_dashboard_stats(data):
    invs = [_invoice("1", 100.0), _invoice("2", 100.0)]
    pays = [{"id": "p1", "customer": "East Port", "amount": 40.0, "applied_invoices": []}]
    stats = data.get_dashboard_stats(invs, pays, [{"name": "East Port"}])
    assert stats["invoice_count"] == 2
    assert stats["total_invoiced"] == pytest.approx(200.0)
    assert stats["total_paid"] == pytest.approx(40.0)
    assert stats["total_outstanding"] == pytest.approx(160.0)
    assert stats["avg_invoice"] == pytest.approx(100.0)


def test_dashboard_stats_empty(data):
    stats = data.get_dashboard_stats([], [], [])
    assert stats["invoice_count"] == 0
    assert stats["avg_invoice"] == 0.0
    assert stats["top_debtors"] == []


# ------------------------------------------------------------------ مرآة إكسل
def _full_invoice(invno="010001", **over):
    inv = {
        "invno": invno, "date": "2026-07-11", "customer": "East Port",
        "declno": "1", "decldate": "", "port": "", "containercount": "",
        "goodstype": "", "origin": "", "notes": "ملاحظة",
        "items": [{"label": "", "labelEn": "", "dinar": 0, "fils": 0}] * 24,
    }
    inv.update(over)
    return inv


def test_excel_mirror_written_after_save(data, tmp_path):
    data.save_customer({"name": "East Port", "company_code": "01"})
    items = [{"label": "خدمة مخصصة", "labelEn": "Custom", "dinar": 10, "fils": 250}]
    items += [{"label": "", "labelEn": "", "dinar": 0, "fils": 0}] * 23
    data.save_invoice(_full_invoice(items=items))
    data.save_payment({"id": "010001", "customer": "East Port", "amount": 5.0,
                       "applied_invoices": ["010001"]})

    wb = load_workbook(tmp_path / "data.xlsx")
    assert set(wb.sheetnames) == {SHEET_INVOICES, SHEET_CUSTOMERS, SHEET_PAYMENTS}
    assert wb[SHEET_CUSTOMERS].max_row == 2
    assert wb[SHEET_PAYMENTS].max_row == 2

    inv_row = [c.value for c in wb[SHEET_INVOICES][2]]
    notes = inv_row[len(INVOICE_HEADERS) - 2]
    # الملاحظات تحتفظ بتسميات البنود المخصّصة موسومة كما في النسخة القديمة
    assert notes.startswith("ملاحظة")
    assert "خدمة مخصصة" in notes
    assert inv_row[-1] == pytest.approx(10.25)


def test_custom_service_labels_survive_roundtrip(data):
    labels = {"label": "بند مخصص", "labelEn": "Custom item"}
    items = [dict(labels, dinar=3, fils=0)]
    items += [{"label": "", "labelEn": "", "dinar": 0, "fils": 0}] * 23
    data.save_invoice(_full_invoice(items=items))

    inv = data.load_invoices()[0]
    assert inv["items"][0]["label"] == "بند مخصص"
    assert inv["items"][0]["labelEn"] == "Custom item"
    # التسميات الافتراضية للباقي تبقى كما هي
    assert inv["items"][1]["label"] == SERVICE_ITEMS[1][0]


def test_excel_read_roundtrip(data, tmp_path):
    """الكتابة إلى إكسل ثم قراءتها تُعيد نفس القيم (بلا تكرار بنود)."""
    items = [{"label": "", "labelEn": "", "dinar": 7, "fils": 500}]
    items += [{"label": "", "labelEn": "", "dinar": 0, "fils": 0}] * 23
    data.save_invoice(_full_invoice(declno="9", port="الشعب", items=items))

    from customs_billing.data import excel_mirror
    invs, _, _ = excel_mirror.read_workbook_data()
    assert len(invs) == 1
    assert invs[0]["invno"] == "010001"
    assert invs[0]["declno"] == "9"
    assert invs[0]["port"] == "الشعب"
    assert invs[0]["notes"] == "ملاحظة"
    assert len(invs[0]["items"]) == 24
    assert invs[0]["total"] == pytest.approx(7.5)


def test_excel_mirror_reflects_deletions(data, tmp_path):
    data.save_invoice(_full_invoice("010001"))
    data.save_invoice(_full_invoice("010002"))
    data.delete_invoice("010001")

    wb = load_workbook(tmp_path / "data.xlsx")
    assert wb[SHEET_INVOICES].max_row == 2  # صف عنوان + فاتورة واحدة فقط
