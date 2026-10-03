# -*- coding: utf-8 -*-
"""
اختبارات طبقة البيانات (بدون واجهة رسومية) تعمل على ملف Excel مؤقت.
تُشغّل بـ:
    python tests/test_data_layer.py
"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# هذه المجموعة مخصصة لمحرك Excel القديم (وضع التوافق) — المتغيرات تُضبط قبل
# الاستيراد لأن الافتراضي الجديد هو SQLite.
SANDBOX = os.path.join(tempfile.gettempdir(), "restaurant_data_tests")
os.makedirs(SANDBOX, exist_ok=True)
os.environ["RESTAURANT_BACKEND"] = "excel"
os.environ["RESTAURANT_DATA_FILE"] = os.path.join(SANDBOX, "restaurant_data.xlsx")

import restaurant_data as rd

rd.BAK_FILE = rd.DATA_FILE + ".bak.xlsx"
if os.path.exists(rd.DATA_FILE):
    os.remove(rd.DATA_FILE)

passed, failed = [], []


def case(label):
    """مُزيّن بسيط: ينفّذ الدالة ويسجل النجاح أو الفشل دون إيقاف بقية الاختبارات."""
    def deco(fn):
        try:
            fn()
            passed.append(label)
            print("  [ OK ] %s" % label)
        except Exception as exc:                 # noqa: BLE001
            failed.append((label, exc))
            print("  [FAIL] %s -> %s: %s" % (label, type(exc).__name__, exc))
        return fn
    return deco


def expect_error(exc_type, fn, *args, **kwargs):
    """يتحقق من أن الاستدعاء يرفع الاستثناء المتوقع."""
    try:
        fn(*args, **kwargs)
    except exc_type:
        return True
    raise AssertionError("لم يُرفع الاستثناء المتوقع %s" % exc_type.__name__)


def assert_eq(actual, expected, label):
    assert actual == expected, "%s: المتوقع %r والفعل %r" % (label, expected, actual)


def assert_last(rows, key, expected, label):
    assert rows, "%s: لا توجد صفوف" % label
    assert_eq(rd.safe_float(rows[-1].get(key)) if isinstance(expected, (int, float))
              else rows[-1].get(key), expected, label)


def assert_count(rows, expected, label):
    assert_eq(len(rows), expected, label)


print("1) إنشاء ملف البيانات")
@case("ensure_workbook يُنشئ كل الجداول المطلوبة")
def _t_workbook():
    rd.ensure_workbook()
    from openpyxl import load_workbook
    wb = load_workbook(rd.DATA_FILE)
    for sheet in rd.SHEETS.values():
        assert sheet in wb.sheetnames, "الجدول الناقص: %s" % sheet


print("2) منع تكرار الأسماء في الجداول التي مفتاحها الاسم")
@case("حفظ عميلين مختلفين")
def _t_customers():
    rd.save_customer({"name": "شركة الخليج", "type": "شركة"})
    rd.save_customer({"name": "مستشفى الأمل", "type": "مستشفى"})
    assert_count(rd.load_customers(), 2, "عدد العملاء")


@case("رفض عميل مكرر")
def _t_dup_customer():
    expect_error(ValueError, rd.save_customer, {"name": "شركة الخليج"})


@case("رفض عميل بلا اسم")
def _t_blank_customer():
    expect_error(ValueError, rd.save_customer, {"name": "   "})


print("3) التحقق من صيغ التواريخ والأشهر")
@case("رفض تاريخ مشتريات غير صحيح")
def _t_bad_purchase_date():
    expect_error(ValueError, rd.save_purchase,
                 {"date": "15/09/2026", "quantity": "1", "unit_cost": "1"})


@case("رفض شهر راتب غير صحيح")
def _t_bad_salary_month():
    expect_error(ValueError, rd.save_salary, {"staff": "علي", "month": "2026/09"})


@case("رفض تاريخ مناسبة غير صحيح")
def _t_bad_event_date():
    expect_error(ValueError, rd.save_event,
                 {"event_date": "20260915", "customer": "شركة الخليج"})


print("4) حساب الإجماليات في طبقة البيانات")
@case("إجمالي المشتريات = الكمية × سعر الوحدة")
def _t_purchase_total():
    rd.save_purchase({"date": "2026-09-10", "type": "لحوم ودواجن", "item": "دجاج",
                      "quantity": "10", "unit": "كجم", "unit_cost": "1.250"})
    assert_last(rd.load_purchases(), "total_cost", 12.5, "إجمالي المشتريات")


@case("إجمالي طلب المؤسسة = عدد الوجبات × السعر")
def _t_order_total():
    rd.save_institutional_order({"date": "2026-09-10", "customer": "شركة الخليج",
                                 "meal_type": "غداء", "quantity": "50", "unit_price": "1.5"})
    assert_last(rd.load_institutional_orders(), "total", 75.0, "إجمالي الطلب")
    assert_last(rd.load_institutional_orders(), "invoiced", "لا", "حالة الفوترة الافتراضية")


@case("إجمالي المناسبة = عدد الأشخاص × سعر الفرد")
def _t_event_total():
    rd.save_event({"event_date": "2026-09-20", "customer": "مستشفى الأمل",
                   "event_type": "فرح / زفاف", "guests_count": "200",
                   "unit_price": "2.5", "deposit": "100"})
    assert_last(rd.load_events(), "total", 500.0, "إجمالي المناسبة")


print("5) الرواتب: منع التسجيل المزدوج لنفس الشهر")
@case("صافي الراتب = الأساسي + المكافآت - السلف - الخصومات")
def _t_net_salary():
    rd.save_salary({"staff": "علي", "month": "2026-09", "base_salary": "300",
                    "bonuses": "20", "advances": "50", "deductions": "10"})
    assert_last(rd.load_salaries(), "net_salary", 260.0, "صافي الراتب")
    assert_last(rd.load_salaries(), "paid", "لا", "حالة الدفع الافتراضية")


@case("حفظ الراتب مرتين لا يُنشئ صفاً مكرراً")
def _t_salary_upsert():
    rd.save_salary({"staff": "علي", "month": "2026-09", "base_salary": "350"})
    assert_count(rd.load_salaries(), 1, "عدد سجلات الرواتب")
    assert_last(rd.load_salaries(), "net_salary", 350.0, "الصافي بعد التحديث")


@case("إعادة الحفظ تحافظ على وسم «مدفوع»")
def _t_salary_paid_kept():
    row = rd.load_salaries()[0]
    row["paid"] = "نعم"
    rd.save_salary(row, editing_id=row["id"])
    rd.save_salary({"staff": "علي", "month": "2026-09", "base_salary": "350"})
    assert_last(rd.load_salaries(), "paid", "نعم", "وسم الدفع")
print("6) الفواتير: الإصدار والفوترة مرة واحدة فقط")
@case("دورة إصدار الفاتورة كاملة")
def _t_invoice_flow():
    orders = rd.load_institutional_orders()
    events = rd.load_events()
    inv = rd.create_invoice_for_customer("شركة الخليج", [orders[0]["id"]], [events[0]["id"]])
    assert_eq(inv["total"], 475.0, "إجمالي الفاتورة (75 + 500 - 100 عربون)")
    assert_eq(len(rd.invoice_items(inv)), 2, "عدد بنود الفاتورة")
    assert_eq(rd.unbilled_institutional_orders_for_customer("شركة الخليج"), [],
              "بعد الفوترة لا تبقى طلبات غير مفوترة")
    assert_eq(rd.unbilled_events_for_customer("مستشفى الأمل"), [],
              "بعد الفوترة لا تبقى مناسبات غير مفوترة")


@case("رفض فاتورة بلا بنود أو بلا عميل")
def _t_invoice_validation():
    expect_error(ValueError, rd.create_invoice_for_customer, "شركة الخليج", [], [])
    expect_error(ValueError, rd.create_invoice_for_customer, "   ",
                 [rd.load_institutional_orders()[0]["id"]], [])


@case("invoice_items تتعامل مع بنود JSON تالفة")
def _t_invoice_items():
    assert_eq(rd.invoice_items({"items_json": "{ليست JSON"}), [], "JSON تالف")
    assert_eq(rd.invoice_items({"items_json": ""}), [], "JSON فارغ")
    assert_eq(rd.invoice_items({}), [], "بدون حقل بنود")


print("7) دقة مقارنة الفترات في التقارير المالية")
@case("month_bounds يحسب آخر يوم في فبراير وديسمبر")
def _t_month_bounds():
    assert_eq(rd.month_bounds("2026-02"), ("2026-02-01", "2026-02-28"), "فبراير")
    assert_eq(rd.month_bounds("2026-12"), ("2026-12-01", "2026-12-31"), "ديسمبر")
    assert_eq(rd.month_bounds("غير صالح"), (None, None), "شهر غير صالح")


@case("الشهر الجزئي يُحتسب بالتقاطع مع الفترة (الإصلاح)")
def _t_summary_mid_month():
    rd.save_monthly_expense({"month": "2026-09", "type": "إيجار المكان", "amount": "500"})
    s = rd.financial_summary("2026-09-15", "2026-09-30")
    assert_eq(s["salaries_cost"], 350.0, "راتب سبتمبر يجب أن يُحتسب")
    assert_eq(s["monthly_expenses_cost"], 500.0, "مصروف سبتمبر الشهري")
    assert_eq(s["purchases_cost"], 0.0, "المشتريات بتاريخ 2026-09-10 خارج الفترة")
    s2 = rd.financial_summary("2026-06-01", "2026-08-31")
    assert_eq(s2["salaries_cost"], 0.0, "راتب سبتمبر خارج فترة يونيو-أغسطس")
    assert_eq(s2["monthly_expenses_cost"], 0.0, "مصروف سبتمبر خارج الفترة")
    s3 = rd.financial_summary()
    assert_eq(s3["revenue"], 475.0, "الإيرادات الكلية")
    assert_eq(s3["operating_expenses"], 862.5, "المصروفات التشغيلية الكلية")
    assert_eq(s3["operating_profit"], -387.5, "صافي الربح التشغيلي")


print("8) الحذف والترقيم")
@case("حذف مصروف تأسيس")
def _t_delete_setup():
    rd.save_setup_expense({"date": "2026-09-01", "type": "تراخيص وسجلات حكومية",
                           "cost": "1200"})
    assert_count(rd.load_setup_expenses(), 1, "قبل الحذف")
    rd.delete_setup_expense(rd.load_setup_expenses()[0]["id"])
    assert_count(rd.load_setup_expenses(), 0, "بعد الحذف")


@case("ترقيم الفواتير يتصاعد بشكل صحيح")
def _t_invoice_numbering():
    assert_eq(rd.next_invoice_no(), "1002", "رقم الفاتورة التالي")
    assert_count(rd.load_invoices(), 1, "عدد الفواتير")


print("9) الأصناف والمنيو")
@case("حفظ صنف ومنع تكرار الاسم وحذفه")
def _t_products():
    rd.save_product({"name": "دجاج مسحب", "category": "أطباق رئيسية", "unit": "حصة / فرد",
                     "cost": "0.800", "sale_price": "1.500"})
    assert_count(rd.load_products(), 1, "عدد الأصناف")
    expect_error(ValueError, rd.save_product, {"name": "دجاج مسحب"})
    expect_error(ValueError, rd.save_product, {"name": "   "})
    rd.delete_product("دجاج مسحب")
    assert_count(rd.load_products(), 0, "بعد الحذف")


print("10) عروض الأسعار والتحويل إلى مناسبة")
@case("حفظ عرض سعر: الإجمالي والحالة الافتراضية والتحقق من التواريخ")
def _t_quotation_save():
    rd.save_quotation({"date": "2026-09-01", "customer": "شركة الخليج",
                       "event_type": "فرح / زفاف", "event_date": "2026-10-01",
                       "guests_count": "200", "unit_price": "2.000"})
    q = rd.load_quotations()[-1]
    assert_eq(q["total"], 400.0, "إجمالي العرض 200 × 2.000")
    assert_eq(q["status"], "معلق", "الحالة الافتراضية")
    expect_error(ValueError, rd.save_quotation, {"date": "01/09/2026", "customer": "شركة الخليج"})
    expect_error(ValueError, rd.save_quotation,
                 {"date": "2026-09-01", "customer": "شركة الخليج", "event_date": "بكرة"})


@case("تحويل العرض إلى مناسبة يمنع التكرار ويحدّث الحالة")
def _t_quotation_convert():
    events_before = len(rd.load_events())
    q = rd.load_quotations()[-1]
    ev = rd.convert_quotation_to_event(q["id"])
    assert_eq(len(rd.load_events()), events_before + 1, "أُنشئت مناسبة جديدة")
    assert_eq(ev["customer"], "شركة الخليج", "عميل المناسبة من العرض")
    assert_eq(rd.safe_float(ev["total"]), 400.0, "إجمالي المناسبة المحوّلة")
    q2 = rd.load_quotations()[-1]
    assert_eq(q2["status"], "مقبول", "حالة العرض بعد التحويل")
    assert_eq(str(q2["converted_id"]), str(ev["id"]), "ربط العرض بالمناسبة")
    expect_error(ValueError, rd.convert_quotation_to_event, q["id"])
    expect_error(ValueError, rd.convert_quotation_to_event, "99999")


print("11) المدفوعات الجزئية وكشف الحساب")
@case("رفض دفعة بلا فاتورة أو بفاتورة غير موجودة أو بمبلغ غير صالح")
def _t_payment_validation():
    expect_error(ValueError, rd.save_payment, {"date": "2026-09-20", "invoice_no": "", "amount": "10"})
    expect_error(ValueError, rd.save_payment, {"date": "2026-09-20", "invoice_no": "0000", "amount": "10"})
    invno = rd.load_invoices()[0]["invoice_no"]
    expect_error(ValueError, rd.save_payment, {"date": "2026-09-20", "invoice_no": invno, "amount": "0"})
    expect_error(ValueError, rd.save_payment, {"date": "20-09-2026", "invoice_no": invno, "amount": "5"})


@case("دفعات جزئية تحدّث المدفوع وحالة التحصيل تلقائياً")
def _t_payment_flow():
    inv = rd.load_invoices()[0]
    invno = inv["invoice_no"]
    total = rd.safe_float(inv["total"])          # 475
    rd.save_payment({"date": "2026-09-20", "invoice_no": invno,
                     "amount": "200", "method": "كي نت"})
    inv = rd.load_invoices()[0]
    assert_eq(rd.safe_float(inv["paid_amount"]), 200.0, "المدفوع بعد دفعة أولى")
    assert_eq(inv["status"], "جزئية", "الحالة بعد دفعة جزئية")
    assert_eq(rd.invoice_balance(inv), 275.0, "المتبقي على الفاتورة")
    rd.save_payment({"date": "2026-09-21", "invoice_no": invno,
                     "amount": str(total - 200), "method": "نقداً", "customer": ""})
    payments = rd.load_payments()
    assert_eq(payments[-1]["customer"], "شركة الخليج", "العميل يُملأ تلقائياً من الفاتورة")
    inv = rd.load_invoices()[0]
    assert_eq(rd.safe_float(inv["paid_amount"]), total, "اكتمال التحصيل")
    assert_eq(inv["status"], "مدفوعة", "الحالة بعد اكتمال التحصيل")
    last_payment = payments[-1]
    rd.delete_payment(last_payment["id"])
    inv = rd.load_invoices()[0]
    assert_eq(rd.safe_float(inv["paid_amount"]), 200.0, "المدفوع بعد حذف الدفعة الأخيرة")
    assert_eq(inv["status"], "جزئية", "الحالة بعد حذف دفعة")


@case("كشف حساب عميل: مدين ودائن ورصيد متحرك")
def _t_statement():
    s = rd.customer_statement("شركة الخليج")
    assert_eq(s["total_invoiced"], 475.0, "إجمالي فواتير العميل")
    assert_eq(s["total_paid"], 200.0, "إجمالي مدفوعات العميل")
    assert_eq(s["balance"], 275.0, "الرصيد المستحق")
    assert s["rows"], "صفوف الكشف غير فارغة"
    assert_eq(s["rows"][-1]["balance"], 275.0, "آخر رصيد متحرك")
    expect_error(ValueError, rd.customer_statement, "   ")
    s2 = rd.customer_statement("عميل بلا حركات")
    assert_eq(s2["rows"], [], "كشف عميل بلا حركات")
    assert_eq(s2["balance"], 0.0, "رصيد صفر لعميل بلا حركات")


print("12) النثريات اليومية")
@case("حفظ نثريات واحتسابها في الملخص المالي ثم حذفها")
def _t_petty_cash():
    expect_error(ValueError, rd.save_petty_cash, {"date": "20-09-2026"})
    rd.save_petty_cash({"date": "2026-09-22", "type": "وقود وتوصيل",
                        "description": "بنزين سيارة التوصيل", "amount": "7.500"})
    s = rd.financial_summary()
    assert_eq(s["petty_cash_cost"], 7.5, "النثريات في الملخص")
    assert_eq(s["operating_expenses"], 870.0, "المصروفات التشغيلية تشمل النثريات")
    rd.delete_petty_cash(rd.load_petty_cash()[0]["id"])
    assert_eq(rd.financial_summary()["petty_cash_cost"], 0.0, "بعد الحذف")


print("13) جدولة التوصيل")
@case("تسليمات اليوم وإسناد سائق وسيارة لطلب مؤسسة")
def _t_deliveries():
    today = rd.date.today().isoformat()
    rd.save_institutional_order({"date": today, "customer": "شركة الخليج",
                                 "meal_type": "عشاء", "frequency": "طلب لمرة واحدة",
                                 "quantity": "40", "unit_price": "1.000",
                                 "delivery_location": "فرع السالمية"})
    rows = [r for r in rd.deliveries_for_date(today) if r["kind"] == "institutional"]
    assert rows, "طلب اليوم في جدولة التوصيل"
    target = rows[-1]
    rd.save_delivery_assignment(target["kind"], target["id"],
                                driver="سالم", vehicle="سيارة 1", delivered=True)
    t2 = [r for r in rd.deliveries_for_date(today) if str(r["id"]) == str(target["id"])][0]
    assert_eq(t2["driver"], "سالم", "حفظ السائق")
    assert_eq(t2["vehicle"], "سيارة 1", "حفظ السيارة")
    assert_eq(t2["delivered"], "نعم", "وسم التسليم")
    expect_error(ValueError, rd.save_delivery_assignment, "institutional", "99999", "س", "ص")
    expect_error(ValueError, rd.save_delivery_assignment, "نوع_خطأ", "1", "س", "ص")


@case("إسناد سائق لمناسبة عبر جدولة التوصيل")
def _t_event_delivery():
    today = rd.date.today().isoformat()
    rd.save_event({"event_date": today, "customer": "عائلة الصباح", "event_type": "عزيمة عائلية",
                   "guests_count": "50", "unit_price": "2.000"})
    ev_rows = [r for r in rd.deliveries_for_date(today) if r["kind"] == "event"]
    assert ev_rows, "مناسبة اليوم في جدولة التوصيل"
    ev = ev_rows[-1]
    rd.save_delivery_assignment(ev["kind"], ev["id"], driver="أحمد", vehicle="سيارة 2")
    t2 = [r for r in rd.deliveries_for_date(today) if str(r["id"]) == str(ev["id"])][0]
    assert_eq(t2["driver"], "أحمد", "سائق المناسبة")
    assert_eq(t2["qty"], "50 فرد", "عرض عدد الأفراد")


print("14) ترقية مخطط ملف قديم بإضافة الأعمدة الجديدة")
@case("إعادة إنشاء أعمدة الفواتير الجديدة تلقائياً دون فقدان بيانات")
def _t_schema_upgrade():
    from openpyxl import load_workbook
    wb = load_workbook(rd.DATA_FILE)
    ws = wb[rd.SHEETS["invoices"]]
    # محاكاة ملف بإصدار قديم: حذف عمودَي المدفوع والحالة
    ws.delete_cols(7, 2)
    wb.save(rd.DATA_FILE)
    rd._schema_checked = False
    rd.ensure_workbook()
    wb2 = load_workbook(rd.DATA_FILE)
    ws2 = wb2[rd.SHEETS["invoices"]]
    headers = [ws2.cell(row=1, column=c).value for c in range(1, ws2.max_column + 1)]
    assert "المدفوع (د.ك)" in headers, "عمود المدفوع رُجّع"
    assert "حالة التحصيل" in headers, "عمود الحالة رُجّع"
    # البيانات القديمة سليمة والأعمدة الجديدة فارغة
    inv = rd.load_invoices()[0]
    assert_eq(rd.safe_float(inv["total"]), 475.0, "إجمالي الفاتورة سليم بعد الترقية")
    # إعادة الحساب تعيد بناء المدفوع والحالة من الدفعات المسجلة
    rd._recalc_invoice_payment(inv["invoice_no"])
    inv = rd.load_invoices()[0]
    assert_eq(rd.safe_float(inv["paid_amount"]), 200.0, "إعادة حساب المدفوع بعد الترقية")
    assert_eq(inv["status"], "جزئية", "الحالة بعد إعادة الحساب")


print()
print("النتيجة: %d ناجح / %d فاشل" % (len(passed), len(failed)))
for label, exc in failed:
    print(" - %s -> %r" % (label, exc))
sys.exit(1 if failed else 0)