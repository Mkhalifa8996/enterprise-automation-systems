# -*- coding: utf-8 -*-
"""
اختبارات محرك SQLite (المرحلة 1): نفس الواجهة العامة لـ restaurant_data
تعمل على قاعدة restaurant.db — CRUD، دورة الفواتير، الرواتب، التقارير،
النسخ الاحتياطية، والتصدير/الاستيراد من/إلى Excel.

    python tests/test_sqlite_backend.py
"""
import json
import os
import shutil
import sys
import tempfile

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE)

# متغيرات البيئة يجب أن تُضبط قبل استيراد restaurant_data
SANDBOX = os.path.join(tempfile.gettempdir(), "restaurant_sqlite_tests")
shutil.rmtree(SANDBOX, ignore_errors=True)
os.makedirs(SANDBOX, exist_ok=True)
os.environ["RESTAURANT_BACKEND"] = "sqlite"
os.environ["RESTAURANT_DB_FILE"] = os.path.join(SANDBOX, "restaurant.db")
os.environ["RESTAURANT_DATA_FILE"] = os.path.join(SANDBOX, "unused.xlsx")

import restaurant_data as rd          # noqa: E402

assert rd.BACKEND == "sqlite", "BACKEND يجب أن يكون sqlite"

passed, failed = [], []


def case(label):
    def deco(fn):
        try:
            fn()
            passed.append(label)
            print("  [ OK ] %s" % label)
        except Exception as exc:                 # noqa: BLE001
            failed.append((label, exc))
            print("  [FAIL] %s -> %s: %s" % (label, type(exc).__name__, exc))
        return fn
    return deco


def assert_true(cond, label):
    assert cond, label


def assert_eq(actual, expected, label):
    assert actual == expected, "%s: المتوقع %r والفعل %r" % (label, expected, actual)


def _audit_rows():
    path = os.path.join(SANDBOX, rd.AUDIT_DIR_NAME,
                        "audit_%s.csv" % rd.date.today().strftime("%Y-%m"))
    with open(path, encoding="utf-8-sig") as fh:
        return fh.read()


INVOICE_TOTAL = {"value": 0.0}   # يُستخدم في اختبار التقارير لاحقاً

print("1) إنشاء قاعدة البيانات والمخطط")


@case("ensure_workbook ينشئ كل جداول النظام في قاعدة SQLite")
def _t_schema():
    rd.ensure_workbook()
    import sqlite3
    conn = sqlite3.connect(os.environ["RESTAURANT_DB_FILE"])
    tables = {r[0] for r in conn.execute(
        "SELECT name FROM sqlite_master WHERE type='table'")}
    conn.close()
    missing = set(rd.SHEETS) - tables
    assert_true(not missing, "جداول ناقصة: %s" % missing)


print("2) العملاء: إضافة ومنع التكرار والحذف")


@case("حفظ عميلين ورفض التكرار وحذف عميل")
def _t_customers():
    rd.save_customer({"name": "شركة الإصلاح", "type": "شركة", "phone": "900"})
    rd.save_customer({"name": "مستشفى العافية", "type": "مستشفى"})
    assert_eq(len(rd.load_customers()), 2, "عدد العملاء")
    try:
        rd.save_customer({"name": "شركة الإصلاح", "type": "شركة"})
        raise AssertionError("لم يُرفض الاسم المكرر")
    except ValueError:
        pass
    rd.delete_customer("مستشفى العافية")
    assert_eq(len(rd.load_customers()), 1, "بعد الحذف")


print("3) المشتريات وطلبات المؤسسات")


@case("ترقيم تلقائي للمعرفات وإجمالي الكمية × السعر")
def _t_purchases():
    rd.save_purchase({"date": "2026-01-10", "type": "لحوم ودواجن",
                      "item": "دجاج", "quantity": "2.5", "unit": "كجم",
                      "unit_cost": "3.0", "total_cost": "", "supplier": "المورد"})
    rows = rd.load_purchases()
    assert_eq(rows[0]["id"], 1, "المعرف التلقائي")
    assert_eq(rd.safe_float(rows[0]["total_cost"]), 7.5, "2.5 × 3.0")
    assert_eq(len(rd.purchases_for_date("2026-01-10")), 1, "مشتريات يوم معين")


@case("حفظ طلب، تعديله بلا تكرار، ثم وسمه مفوتراً")
def _t_orders():
    rd.save_institutional_order({"date": "2026-02-01", "customer": "شركة الإصلاح",
                                 "meal_type": "غداء", "frequency": "يومي",
                                 "quantity": "3", "unit_price": "2.5",
                                 "delivery_location": "المقر"})
    rows = rd.load_institutional_orders()
    assert_eq(len(rows), 1, "صف واحد")
    assert_eq(rd.safe_float(rows[0]["total"]), 7.5, "3 × 2.5")
    oid = rows[0]["id"]
    rd.save_institutional_order({"date": "2026-02-01", "customer": "شركة الإصلاح",
                                 "meal_type": "غداء", "quantity": "4",
                                 "unit_price": "2.5"}, editing_id=oid)
    rows = rd.load_institutional_orders()
    assert_eq(len(rows), 1, "التعديل لا يُنشئ صفاً جديداً")
    assert_eq(rd.safe_float(rows[0]["total"]), 10.0, "الإجمالي بعد التعديل")
    rd.mark_institutional_order_invoiced(oid, "1001")
    rows = rd.load_institutional_orders()
    assert_eq(rows[0]["invoiced"], "نعم", "وسم مفوتر")
    assert_eq(rows[0]["invoice_no"], "1001", "رقم الفاتورة في الطلب")


print("4) المناسبات")


@case("حفظ مناسبة بإجمالي صحيح وupcoming_events")
def _t_events():
    rd.save_event({"event_date": "2026-03-05", "customer": "شركة الإصلاح",
                   "event_type": "تخرج", "guests_count": "20",
                   "unit_price": "4.0", "deposit": "5.0", "location": "قاعة"})
    events = rd.load_events()
    assert_eq(rd.safe_float(events[0]["total"]), 80.0, "20 × 4.0")
    assert_eq(len(rd.upcoming_events("2026-01-01")), 1, "مناسبة قادمة")
    assert_eq(len(rd.upcoming_events("2026-04-01")), 0, "فائتة مستبعدة")


print("5) دورة الفاتورة الكاملة")


@case("فاتورة تجمع طلباً ومناسبة، وتوسم الاثنين")
def _t_invoice_cycle():
    order = rd.load_institutional_orders()[0]
    event = rd.load_events()[0]
    invoice = rd.create_invoice_for_customer("شركة الإصلاح",
                                             [order["id"]], [event["id"]])
    assert_eq(invoice["invoice_no"], "1001", "أول رقم فاتورة")
    items = rd.invoice_items(invoice)
    assert_eq(len(items), 2, "بنود الفاتورة")
    saved = [i for i in rd.load_invoices()
             if str(i["invoice_no"]) == str(invoice["invoice_no"])]
    assert_eq(len(saved), 1, "الفاتورة محفوظة مرة واحدة")
    # الإجمالي = طلب 10.0 + (مناسبة 80.0 - عربون 5.0)
    assert_eq(rd.safe_float(invoice["total"]), 85.0, "10 + 80 - 5")
    INVOICE_TOTAL["value"] = rd.safe_float(invoice["total"])
    assert_eq(rd.unbilled_institutional_orders_for_customer("شركة الإصلاح"), [],
              "لا طلبات غير مفوترة")
    assert_eq(rd.unbilled_events_for_customer("شركة الإصلاح"), [],
              "لا مناسبات غير مفوترة")
    assert_eq(rd.next_invoice_no(), "1002", "الرقم التالي")


@case("رفض فاتورة بلا عميل أو بلا بنود صالحة")
def _t_invoice_rejects():
    try:
        rd.create_invoice_for_customer("", [], [])
        raise AssertionError("لم يُرفض الحفظ بلا عميل")
    except ValueError:
        pass
    try:
        rd.create_invoice_for_customer("شركة الإصلاح", [999], [999])
        raise AssertionError("لم يُرفض الحفظ بلا بنود صالحة")
    except ValueError:
        pass


print("6) الرواتب: upsert شهري ورفض صيغة خاطئة")


@case("حفظ الراتب مرتين لنفس الشهر لا يُنشئ صفاً مكرراً ويحفظ وسم الدفع")
def _t_salaries():
    rd.save_staff({"name": "أحمد", "job_title": "طباخ مساعد",
                   "base_salary": "300.0"})
    rd.save_salary({"staff": "أحمد", "month": "2026-02", "base_salary": "300",
                    "bonuses": "10", "advances": "5", "deductions": "0",
                    "paid": "نعم"})
    rd.save_salary({"staff": "أحمد", "month": "2026-02", "base_salary": "300",
                    "bonuses": "0", "advances": "0", "deductions": "10"})
    rows = rd.load_salaries()
    assert_eq(len(rows), 1, "صف واحد لنفس الشهر")
    assert_eq(rd.safe_float(rows[0]["net_salary"]), 290.0, "300+0-0-10")
    assert_eq(rows[0]["paid"], "نعم", "وسم الدفع السابق محفوظ")
    try:
        rd.save_salary({"staff": "أحمد", "month": "فبراير"})
        raise AssertionError("لم يُرفض شهر بصيغة خاطئة")
    except ValueError:
        pass


print("7) التقارير المالية على محرك SQLite")


@case("financial_summary يحسب الإيرادات والمصروفات على كامل الفترات")
def _t_summary():
    rd.save_monthly_expense({"month": "2026-03", "type": "إيجار المكان",
                             "amount": "150.0"})
    summary = rd.financial_summary()          # بلا حدود: كل السجلات
    assert_eq(summary["revenue"], INVOICE_TOTAL["value"], "الإيرادات = الفاتورة")
    assert_eq(summary["purchases_cost"], 7.5, "المشتريات")
    assert_eq(summary["salaries_cost"], 290.0, "راتب فبراير")
    assert_eq(summary["monthly_expenses_cost"], 150.0, "مصروف شهر مارس")
    assert_eq(summary["operating_expenses"], 447.5, "7.5+290+150")
    assert_eq(summary["operating_profit"],
              round(INVOICE_TOTAL["value"] - 447.5, 3), "الربح التشغيلي")
    assert_true("setup_expenses_total" in summary, "مفاتيح الملخص كاملة")


@case("financial_summary يستبعد ما خارج الفترة المحددة")
def _t_summary_period():
    summary = rd.financial_summary("2026-03-01", "2026-03-31")
    assert_eq(summary["revenue"], 0.0, "الفاتورة (اليوم) خارج فترة مارس")
    assert_eq(summary["purchases_cost"], 0.0, "مشتريات يناير مستبعدة")
    assert_eq(summary["salaries_cost"], 0.0, "راتب فبراير مستبعد")
    assert_eq(summary["monthly_expenses_cost"], 150.0, "مصروف مارس داخل الفترة")


print("8) النسخ الاحتياطية للقاعدة")


@case("نسخة يومية من ملف قاعدة البيانات داخل backups/")
def _t_backup():
    path = rd.backup_data_file(force=True)
    assert_true(path and os.path.exists(path), "النسخة أُنشئت")
    assert_eq(rd.backup_data_file(), None, "لا تكرار في نفس اليوم")
    files = os.listdir(os.path.join(SANDBOX, "backups"))
    assert_true(any(f.startswith("restaurant_db_") for f in files),
                "تسمية نسخ القاعدة")


print("9) التصدير إلى Excel والاستيراد منه (الترحيل)")


@case("export_excel ينشئ ملفاً بأوراق مطابقة وimport_excel يقرأه في قاعدة جديدة")
def _t_roundtrip():
    exported = os.path.join(SANDBOX, "export_test.xlsx")
    rd.export_excel(exported)
    from openpyxl import load_workbook
    wb = load_workbook(exported)
    assert_true(rd.SHEETS["customers"] in wb.sheetnames, "أوراق بنفس الأسماء")
    counts_before = {k: len(rd._db.load_rows(k)) for k in rd.SHEETS}
    # قاعدة ثانية فارغة ثم استيراد الملف المصدر إليها
    rd._db.DB_FILE = os.path.join(SANDBOX, "imported.db")
    rd._db.ensure()
    counts = rd.import_excel(exported)
    assert_eq(counts.get("customers"), counts_before["customers"],
              "عدد العملاء المستورد")
    assert_eq(counts.get("invoices"), counts_before["invoices"],
              "عدد الفواتير المستورد")
    rows = rd.load_customers()
    assert_true(any(c["name"] == "شركة الإصلاح" for c in rows),
                "بيانات العميل سليمة بعد الترحيل")
    inv = rd.load_invoices()[0]
    assert_eq(rd.safe_float(inv["total"]), INVOICE_TOTAL["value"],
              "إجمالي الفاتورة بعد الترحيل")
    rd._db.DB_FILE = os.environ["RESTAURANT_DB_FILE"]   # العودة للقاعدة الأصلية


print("9.5) الميزات الجديدة على محرك SQLite")


@case("عرض سعر وأصناف ونثريات تعمل على SQLite مثل Excel")
def _t_new_features():
    # تعمل في نهاية الملف حتى لا تغيّر ترتيب الصفوف الذي تعتمده الأقسام السابقة
    rd.save_quotation({"date": "2026-01-05", "customer": "شركة الإصلاح",
                       "event_type": "مناسبة رسمية / شركة", "event_date": "2026-02-01",
                       "guests_count": "60", "unit_price": "2.000"})
    q = rd.load_quotations()[-1]
    assert_eq(rd.safe_float(q["total"]), 120.0, "إجمالي العرض على SQLite")
    ev = rd.convert_quotation_to_event(q["id"])
    q2 = rd.load_quotations()[-1]
    assert_eq(q2["status"], "مقبول", "حالة العرض بعد التحويل")
    assert_eq(str(ev["customer"]), "شركة الإصلاح", "مناسبة محوّلة على SQLite")
    rd.save_petty_cash({"date": "2026-01-06", "type": "وقود وتوصيل", "amount": "3.500"})
    summary = rd.financial_summary()
    assert_eq(summary["petty_cash_cost"], 3.5, "النثريات في ملخص SQLite")
    rd.save_product({"name": "أرز بسمتي", "category": "أطباق رئيسية",
                     "cost": "0.500", "sale_price": "1.000"})
    assert_true(len(rd.load_products()) == 1, "الأصناف على SQLite")


print("9.6) الاستيراد التلقائي عند أول إقلاع")


@case("قاعدة فارغة + ملف Excel بجانبها تُستورد تلقائياً مع نسخة احتياطية")
def _t_auto_import():
    import shutil
    import subprocess
    import sys
    sandbox = os.path.join(SANDBOX, "auto_import")
    shutil.rmtree(sandbox, ignore_errors=True)
    os.makedirs(sandbox, exist_ok=True)
    excel_path = os.path.join(sandbox, "restaurant_data.xlsx")
    db_path = os.path.join(sandbox, "restaurant.db")
    # 1) بناء ملف Excel قديم حقيقي في عملية فرعية على محرك Excel
    seed = (
        "import os, sys; sys.path.insert(0, %r); "
        "os.environ['RESTAURANT_BACKEND']='excel'; "
        "os.environ['RESTAURANT_DATA_FILE']=%r; "
        "import restaurant_data as rd; "
        "rd.save_customer({'name': 'عميل الترحيل', 'type': 'شركة'}); "
        "assert len(rd.load_customers()) == 1"
        % (BASE, excel_path))
    r1 = subprocess.run([sys.executable, "-c", seed], capture_output=True,
                        text=True, cwd=BASE, timeout=120)
    assert_true(r1.returncode == 0, "بناء ملف Excel القديم: %s" % r1.stderr[-500:])
    # 2) أول إقلاع على SQLite مع وجود Excel: استيراد تلقائي عبر ensure_workbook
    boot = (
        "import os, sys; sys.path.insert(0, %r); "
        "os.environ['RESTAURANT_BACKEND']='sqlite'; "
        "os.environ['RESTAURANT_DB_FILE']=%r; "
        "os.environ['RESTAURANT_DATA_FILE']=%r; "
        "import restaurant_data as rd; "
        "rd.ensure_workbook(); "
        "names=[c['name'] for c in rd.load_customers()]; "
        "assert 'عميل الترحيل' in names, names; "
        "print('imported-ok')"
        % (BASE, db_path, excel_path))
    r2 = subprocess.run([sys.executable, "-c", boot], capture_output=True,
                        text=True, cwd=BASE, timeout=120)
    assert_true(r2.returncode == 0 and "imported-ok" in r2.stdout,
                "الإقلاع الأول استورد Excel: %s" % r2.stderr[-500:])
    backups = os.listdir(os.path.join(sandbox, rd.BACKUP_DIR_NAME))
    assert_true(any(n.endswith(".xlsx") for n in backups),
                "نسخة احتياطية من Excel قبل الاستيراد")
    # 3) لا يتكرر: كتابة صف جديد ثم إقلاع ثانٍ لا يعيد الاستيراد ولا يمس الصف
    boot2 = boot.replace("print('imported-ok')",
                         "rd.save_customer({'name': 'عميل جديد', 'type': 'شركة'}); "
                         "rd.ensure_workbook(); "
                         "names=[c['name'] for c in rd.load_customers()]; "
                         "assert sorted(names)==sorted(['عميل الترحيل', 'عميل جديد']), names; "
                         "print('once-ok')")
    r3 = subprocess.run([sys.executable, "-c", boot2], capture_output=True,
                        text=True, cwd=BASE, timeout=120)
    assert_true(r3.returncode == 0 and "once-ok" in r3.stdout,
                "الاستيراد مرة واحدة فقط: %s" % r3.stderr[-500:])


print("10) سجل الحركات وقفل النسخة الواحدة")


def _lock_is_held(path):
    try:
        handle = open(path, "a+")
        if os.name == "nt":
            import msvcrt
            msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
            msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
        handle.close()
        return False
    except OSError:
        return True


@case("log_action يكتب CSV وacquire/release القفل يعملان على مسار القاعدة")
def _t_audit_lock():
    text = _audit_rows()
    assert_true("إضافة" in text and "حذف" in text and "إصدار فاتورة" in text,
                "السجل يوثّق الإضافة والحذف وإصدار الفاتورة")
    assert_true(rd.acquire_instance_lock(), "اكتساب القفل")
    assert_true(rd.acquire_instance_lock(), "القفل نفسه لا يُكتسب مرتين (مقبض واحد)")
    rd.release_instance_lock()
    lock_file = os.environ["RESTAURANT_DB_FILE"] + ".lock"
    assert_true(not _lock_is_held(lock_file), "القفل مُحرَّر بعد الإغلاق")


print()
print("النتيجة: %d ناجح / %d فاشل" % (len(passed), len(failed)))
if failed:
    for label, exc in failed:
        print("  فاشل: %s -> %s" % (label, exc))
    print("اختبارات محرك SQLite فشلت.")
    sys.exit(1)
print("كل اختبارات محرك SQLite نجحت.")
