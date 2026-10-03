# -*- coding: utf-8 -*-
"""
أداة الترحيل من Excel إلى SQLite — migrate_to_sqlite.py (يوم التحويل)

ما تفعله بالترتيب:
  1) تأخذ نسخة احتياطية مؤرَّخة من ملف Excel قبل أي شيء.
  2) تجهّز قاعدة SQLite (افتراضياً restaurant.db بجوار البرنامج).
  3) تستورد كل الأوراق من Excel إلى القاعدة (داخل معاملات، وبنفس أسماء
     الأوراق وترتيب الأعمدة).
  4) تتحقق: تقارن عدد الصفوف في كل جدول بين المصدر والهدف.
  5) تطبع الخطوة التالية لتشغيل البرنامج على القاعدة.

ملف Excel لا يُفتح للكتابة إطلاقاً — يبقى مرجعاً سليماً يمكن العودة إليه
في أي وقت عبر RESTAURANT_BACKEND=excel.

الاستخدام:
  python migrate_to_sqlite.py                                  # المسارات الافتراضية
  python migrate_to_sqlite.py --excel ملف.xlsx --db ملف.db --yes
"""
import argparse
import os
import shutil
import sys
from datetime import datetime

# هذه الأداة جوهرها محرك SQLite؛ نضمن تفعيله قبل استيراد restaurant_data
# حتى تُوصَل طبقة القاعدة بوصف الجداول ودالة السجل تلقائياً.
os.environ["RESTAURANT_BACKEND"] = "sqlite"

import restaurant_data as rd          # noqa: E402
from core.data import sqlite as rdb     # noqa: E402
from openpyxl import load_workbook    # noqa: E402


def count_excel_rows(path):
    """عدد الصفوف غير الفارغة في كل ورقة (بدون صف العناوين)."""
    wb = load_workbook(path, data_only=True)
    counts = {}
    for sheet_key, sheet_name in rd.SHEETS.items():
        if sheet_name not in wb.sheetnames:
            continue
        n = 0
        for row in wb[sheet_name].iter_rows(min_row=2, values_only=True):
            if any(v not in (None, "") for v in row):
                n += 1
        counts[sheet_key] = n
    return counts


def main():
    parser = argparse.ArgumentParser(
        description="ترحيل بيانات المطعم من Excel إلى SQLite")
    parser.add_argument("--excel", default=rd.DATA_FILE,
                        help="ملف Excel المصدر (افتراضياً restaurant_data.xlsx)")
    parser.add_argument("--db", default=rdb.DB_FILE,
                        help="ملف قاعدة SQLite الهدف (افتراضياً restaurant.db)")
    parser.add_argument("--yes", action="store_true",
                        help="بدون سؤال تأكيد (للاستخدام الآلي)")
    args = parser.parse_args()

    if not os.path.exists(args.excel):
        print("خطأ: ملف Excel غير موجود: %s" % args.excel)
        return 1
    if os.path.exists(args.db) and not args.yes:
        answer = input("القاعدة %s موجودة وسيُستبدل محتواها بالمستورد. متابعة؟ (y/n) "
                       % args.db)
        if answer.strip().lower() not in ("y", "yes"):
            print("أُلغي الترحيل.")
            return 1

    # 1) نسخة احتياطية مؤرَّخة من Excel داخل مجلد النسخ الاحتياطية
    backup_dir = os.path.join(os.path.dirname(os.path.abspath(args.excel)),
                              rd.BACKUP_DIR_NAME)
    os.makedirs(backup_dir, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    base = os.path.splitext(os.path.basename(args.excel))[0]
    backup = os.path.join(backup_dir, "%s_pre_migration_%s.xlsx" % (base, stamp))
    shutil.copy2(args.excel, backup)
    print("1) نسخة احتياطية من Excel: %s" % backup)

    # 2) تجهيز القاعدة على المسار المطلوب
    rdb.DB_FILE = args.db
    try:
        rdb.ensure()
    except Exception as exc:
        print("فشل تجهيز قاعدة البيانات: %s" % exc)
        print("تأكد أن المسار قابل للكتابة وأن البرنامج غير مفتوح على نفس القاعدة.")
        return 1
    print("2) قاعدة البيانات جاهزة: %s" % args.db)

    # 3) الاستيراد
    try:
        counts = rdb.import_excel(args.excel)
    except Exception as exc:
        print("فشل الاستيراد: %s" % exc)
        print("إن كان البرنامج مفتوحاً على نفس القاعدة فأغلقه ثم أعد المحاولة.")
        return 1
    print("3) استُورد %d صف في %d جدول" % (sum(counts.values()), len(counts)))

    # 4) التحقق: مقارنة عدد الصفوف لكل جدول
    src = count_excel_rows(args.excel)
    problems = []
    print("4) التحقق (المصدر -> القاعدة):")
    for key, sheet_name in rd.SHEETS.items():
        s = src.get(key, 0)
        d = rdb.count_rows(key)
        ok = (s == d)
        if not ok:
            problems.append(key)
        print("   [%s] %s: %d -> %d"
              % ("مطابق " if ok else "مختلف", sheet_name, s, d))

    # 5) الخطوة التالية
    print("5) انتهى. شغّل البرنامج عادياً — SQLite هو الافتراضي الآن:")
    print("     python restaurant_desktop_app.py")
    print("   (وضع Excel القديم متاح فقط عبر: set RESTAURANT_BACKEND=excel)")
    if problems:
        print("تحذير: جداول لم يتطابق عددها: %s" % ", ".join(problems))
        return 1
    print("كل الجداول مطابقة — الترحيل نجح.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
