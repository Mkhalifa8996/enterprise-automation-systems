# -*- coding: utf-8 -*-
"""
نظام فواتير - شركة الاستدامة لتقديم الخدمات اللوجستية
طبقة الوصول إلى البيانات (لا تعتمد على tkinter) — تحتوي كل دوال
القراءة والكتابة المباشرة على ملف data.xlsx.
"""

import json
import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

# ------------------------------------------------------------------
# الإعدادات العامة
# ------------------------------------------------------------------
APP_DIR = os.path.dirname(os.path.abspath(__file__))
FROZEN = getattr(sys, "frozen", False)

if FROZEN:
    BASE_DIR = os.path.dirname(os.path.abspath(sys.argv[0]))
    RESOURCE_DIR = getattr(sys, "_MEIPASS", APP_DIR)
else:
    BASE_DIR = APP_DIR
    RESOURCE_DIR = APP_DIR

DATA_FILE = os.path.join(BASE_DIR, "data.xlsx")
ICON_FILE = os.path.join(RESOURCE_DIR, "icon.png")   # أيقونة التطبيق (شعار مبسط بدون نص)
LOGO_FILE = os.path.join(RESOURCE_DIR, "logo.png")   # الشعار الكامل مع اسم الشركة

SHEET_INVOICES = "الفواتير"
SHEET_CUSTOMERS = "العملاء"
SHEET_PAYMENTS = "الدفعات"

FONT_NAME = "Segoe UI"  # يدعم العربية بشكل جيد على ويندوز؛ Tk سيستخدم بديلاً مناسباً تلقائياً على ماك/لينكس


def get_tk_font(size=10, weight="normal"):
    """يعيد صيغة خط متوافقة مع Tkinter على ويندوز/لينكس/ماك."""
    return (FONT_NAME, size, weight)


# 24 بند خدمة كما في الفاتورة الورقية الأصلية (عربي، إنجليزي)
SERVICE_ITEMS = [
    ("إذن تسليم بوليصة الشحن", "DELIVERY ORDER SHIPPING"),
    ("أجور استلام إذن التسليم", "REVIEWING CHARGE - DELIVERY ORDER"),
    ("أجور الشحن والخدمات اللوجستية", "FREIGHT CHARGE SHIPPING LOGISTIC"),
    ("أرضيات الوكيل او محطة الحاويات", "FREIGHT FORWARDWE FLOORING - FEES"),
    ("رسوم تصديقات الفاتورة والمنشأ", "LEGALIZATION FEES INVOICE"),
    ("أجور تصديقات الفواتير", "LEGALIZAION CHARGES"),
    ("رسوم جلوبال", "GLOBAL FEES"),
    ("الرسوم الجمركية (الجمارك)", "CUSTOMS DUTY"),
    ("عمولة البنك للرسوم الجمركية", "BANK COMMISSION"),
    ("خدمة الإعفاء الجمركي", "EXEMPTION FROM LNDUSTRAY"),
    ("رسوم أفراج الصناعة أو البيئة", "APPROVAL OF LNDUSTRAY"),
    ("رسوم خدمات الأفراج", "APPROVAL OF LNDUSTRIAL - FEES"),
    ("رسوم خدمات الميناء", "PORT SERVICE - FEES"),
    ("أجور كشف جمركي", "CUSTOMS CHARGE"),
    ("أجور خدمات النقل", "TRANSPORTATION CHARGE"),
    ("أجور خدمات العمال", "SERVICE - FEES"),
    ("أجور الكرين أو الرافعة", "FORKLIFT AND CRANE - FEES"),
    ("غرامة طبالي", "FINE FOR DRUMS"),
    ("أجور البيان الجمركي", "CUSTOMS - FEES"),
    ("رسوم خدمات جمركية سحب إسكان", "CUSTOMS SERVICE - FEES"),
    ("أجور التخليص والأتعاب", "CLEARANCE FEES CHARGE"),
    ("طباعة وتصوير المستندات", "PRINTING AND COPYING - FEES"),
    ("أجور الاستلام والمتابعة", "RECEIVING & FOLLOW UP SERVICE"),
    ("أجور خدمات أخرى", "OTHER SERVICE CHARGES"),
]

# الحقول الأساسية + تفاصيل الشحنة (المفتاح، التسمية المعروضة)
META_FIELDS = [
    ("invno", "رقم الفاتورة"),
    ("date", "التاريخ"),
    ("customer", "السادة / اسم العميل"),
    ("declno", "رقم البيان الجمركي"),
    ("decldate", "تاريخ البيان الجمركي"),
    ("port", "المنفذ"),
    ("containercount", "عدد الحاويات"),
    ("goodstype", "نوع البضاعة"),
    ("origin", "بلد المنشأ"),
]

INVOICE_HEADERS = [lbl for _, lbl in META_FIELDS]
for ar, en in SERVICE_ITEMS:
    INVOICE_HEADERS.append(ar + " - دينار")
    INVOICE_HEADERS.append(ar + " - فلس")
INVOICE_HEADERS += ["ملاحظات", "الإجمالي"]

CUSTOMER_HEADERS = ["اسم العميل", "الهاتف", "العنوان", "ملاحظات", "رمز الشركة"]

PAYMENT_HEADERS = ["رقم السند", "اسم العميل", "التاريخ", "المبلغ (د.ك)", "ملاحظات", "الفواتير المسددة"]

HEADER_FILL = PatternFill(start_color="0F2B46", end_color="0F2B46", fill_type="solid")
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")


class SaveError(Exception):
    """يُطلق عند تعذر الحفظ على data.xlsx (غالباً لأن الملف مفتوح في برنامج آخر)."""
    pass


def _safe_save(wb):
    """يحفظ ملف العمل، ويحوّل أخطاء القرص الشائعة إلى رسالة عربية واضحة."""
    try:
        wb.save(DATA_FILE)
    except PermissionError:
        raise SaveError(
            "تعذر الحفظ في ملف data.xlsx.\n"
            "الرجاء إغلاق الملف إذا كان مفتوحاً في برنامج إكسل ثم المحاولة مرة أخرى."
        )
    except OSError as e:
        raise SaveError(f"تعذر الحفظ في ملف data.xlsx:\n{e}")


# ------------------------------------------------------------------
# طبقة الوصول إلى ملف الإكسل (كل قراءة/كتابة تتم مباشرة على data.xlsx)
# ------------------------------------------------------------------
def ensure_workbook():
    """ينشئ ملف data.xlsx مع أوراق العمل والعناوين إذا لم يكن موجوداً،
    ويضيف ورقة الدفعات تلقائياً لملفات data.xlsx القديمة التي لا تحتوي عليها."""
    if not os.path.exists(DATA_FILE):
        wb = Workbook()
        ws1 = wb.active
        ws1.title = SHEET_INVOICES
        _write_header(ws1, INVOICE_HEADERS)
        ws2 = wb.create_sheet(SHEET_CUSTOMERS)
        _write_header(ws2, CUSTOMER_HEADERS)
        ws3 = wb.create_sheet(SHEET_PAYMENTS)
        _write_header(ws3, PAYMENT_HEADERS)
        _safe_save(wb)
        return

    # الملف موجود مسبقاً: تأكد من وجود ورقة الدفعات (توافقاً مع الملفات المنشأة قبل هذه الميزة)
    wb = load_workbook(DATA_FILE)
    changed = False
    if SHEET_PAYMENTS not in wb.sheetnames:
        ws3 = wb.create_sheet(SHEET_PAYMENTS)
        _write_header(ws3, PAYMENT_HEADERS)
        changed = True

    # توافقاً مع ملفات الدفعات القديمة التي لا تحتوي عمود "الفواتير المسددة"
    if SHEET_PAYMENTS in wb.sheetnames:
        wsp = wb[SHEET_PAYMENTS]
        applied_col = len(PAYMENT_HEADERS)  # عمود "الفواتير المسددة"
        if wsp.cell(row=1, column=applied_col).value != PAYMENT_HEADERS[applied_col - 1]:
            c = wsp.cell(row=1, column=applied_col, value=PAYMENT_HEADERS[applied_col - 1])
            c.font = HEADER_FONT
            c.fill = HEADER_FILL
            c.alignment = Alignment(horizontal="center", vertical="center")
            changed = True

    # توافقاً مع ملفات العملاء القديمة التي لا تحتوي عمود "رمز الشركة"
    if SHEET_CUSTOMERS in wb.sheetnames:
        ws2 = wb[SHEET_CUSTOMERS]
        code_col = len(CUSTOMER_HEADERS)  # عمود "رمز الشركة"
        if ws2.cell(row=1, column=code_col).value != CUSTOMER_HEADERS[code_col - 1]:
            c = ws2.cell(row=1, column=code_col, value=CUSTOMER_HEADERS[code_col - 1])
            c.font = HEADER_FONT
            c.fill = HEADER_FILL
            c.alignment = Alignment(horizontal="center", vertical="center")
            changed = True

    if changed:
        _safe_save(wb)


def _write_header(ws, headers):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"


def load_invoices():
    """يقرأ جميع الفواتير من الملف ويعيدها كقائمة قواميس."""
    ensure_workbook()
    wb = load_workbook(DATA_FILE)
    ws = wb[SHEET_INVOICES]
    rows = []
    for r in range(2, ws.max_row + 1):
        first_cell = ws.cell(row=r, column=1).value
        if first_cell in (None, ""):
            continue
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, len(INVOICE_HEADERS) + 1)]
        rows.append(row_to_invoice_dict(row_vals))
    return rows


def row_to_invoice_dict(row_vals):
    d = {}
    for i, (key, _) in enumerate(META_FIELDS):
        d[key] = row_vals[i] if row_vals[i] is not None else ""
    base = len(META_FIELDS)
    items = []
    labels_data = None
    notes_idx = base + len(SERVICE_ITEMS) * 2
    if len(row_vals) > notes_idx:
        notes_raw = row_vals[notes_idx]
        notes_value, labels_data = _split_service_labels_from_notes(notes_raw)
        d["notes"] = notes_value if notes_value is not None else ""
    else:
        d["notes"] = ""
    for idx, (ar, en) in enumerate(SERVICE_ITEMS):
        dinar = row_vals[base + idx * 2] if base + idx * 2 < len(row_vals) else ""
        fils = row_vals[base + idx * 2 + 1] if base + idx * 2 + 1 < len(row_vals) else ""
        item_label = ar
        item_label_en = en
        if labels_data and idx < len(labels_data):
            item_label = labels_data[idx].get("ar", ar)
            item_label_en = labels_data[idx].get("en", en)
        items.append({"label": item_label, "labelEn": item_label_en,
                      "dinar": dinar if dinar is not None else "",
                      "fils": fils if fils is not None else ""})
    d["items"] = items
    d["total"] = row_vals[notes_idx + 1] if notes_idx + 1 < len(row_vals) else 0
    return d


def invoice_dict_to_row(d):
    row = [d.get(key, "") for key, _ in META_FIELDS]
    total_fils = 0
    for it in d["items"]:
        dinar = safe_int(it.get("dinar"))
        fils = safe_int(it.get("fils"))
        row.append(dinar if dinar else "")
        row.append(fils if fils else "")
        total_fils += dinar * 1000 + fils
    notes_value = d.get("notes", "")
    encoded_labels = _encode_service_labels_for_notes(d.get("items") or [])
    if encoded_labels:
        notes_value = f"{notes_value}\n__SERVICE_LABELS__{encoded_labels}__" if notes_value else f"__SERVICE_LABELS__{encoded_labels}__"
    row.append(notes_value)
    row.append(round(total_fils / 1000, 3))
    return row


def safe_int(v):
    try:
        if v in (None, ""):
            return 0
        return int(round(float(v)))
    except (ValueError, TypeError):
        return 0


def safe_float(v):
    try:
        if v in (None, ""):
            return 0.0
        return float(v)
    except (ValueError, TypeError):
        return 0.0


def _default_service_labels_payload():
    return json.dumps([{"ar": ar, "en": en} for ar, en in SERVICE_ITEMS], ensure_ascii=False)


def _encode_service_labels(items):
    """يُنشئ نسخة JSON من تسميات الخدمات الحالية."""
    labels = []
    for idx, (ar, en) in enumerate(SERVICE_ITEMS):
        item = items[idx] if items and idx < len(items) else {}
        ar_label = item.get("label") if item else None
        en_label = item.get("labelEn") if item else None
        labels.append({
            "ar": ar_label if ar_label not in (None, "") else ar,
            "en": en_label if en_label not in (None, "") else en,
        })
    return json.dumps(labels, ensure_ascii=False)


def _encode_service_labels_for_notes(items):
    payload = _encode_service_labels(items)
    if payload == _default_service_labels_payload():
        return ""
    return payload


def _decode_service_labels(value):
    if not isinstance(value, str):
        return None
    marker = "__SERVICE_LABELS__"
    if value.startswith(marker):
        value = value[len(marker):]
    try:
        decoded = json.loads(value)
    except (TypeError, ValueError):
        return None
    return decoded if isinstance(decoded, list) else None


def _split_service_labels_from_notes(notes_value):
    if not isinstance(notes_value, str):
        return notes_value, None
    marker = "__SERVICE_LABELS__"
    if marker not in notes_value:
        return notes_value, None
    parts = notes_value.split(marker, 1)
    if len(parts) != 2:
        return notes_value, None
    notes_text = parts[0].rstrip()
    payload = parts[1]
    if payload.endswith("__"):
        payload = payload[:-2]
    return notes_text, _decode_service_labels(payload)


def normalize_company_code(v):
    """يحوّل رمز الشركة إلى رقمين (01-99) كنص، أو '' إن كان غير صالح/فارغ."""
    if v in (None, ""):
        return ""
    try:
        n = int(str(v).strip())
    except (ValueError, TypeError):
        return ""
    if n < 1 or n > 99:
        return ""
    return f"{n:02d}"


def find_customer_by_company_code(company_code, exclude_name=None):
    """يعيد اسم أول عميل آخر يستخدم نفس رمز الشركة (لمنع التكرار)، أو None إن لم يوجد."""
    code = normalize_company_code(company_code)
    if not code:
        return None
    for c in load_customers():
        if c["company_code"] == code and c["name"] != exclude_name:
            return c["name"]
    return None


def next_available_company_code():
    """يعيد أصغر رمز شركة غير مستخدم بعد (من 01 إلى 99) لتعيينه تلقائياً لعميل جديد."""
    used = {c["company_code"] for c in load_customers() if c["company_code"]}
    for n in range(1, 100):
        code = f"{n:02d}"
        if code not in used:
            return code
    raise ValueError("تم الوصول للحد الأقصى لعدد الشركات (99)؛ لا يوجد رمز شركة متاح.")


def next_invoice_number(company_code):
    """يولّد رقم الفاتورة التالي لعميل معين بصيغة: رمز الشركة (رقمان) + رقم تسلسلي (4 أرقام)
    مثال: الرمز 25 مع الفاتورة الثانية لهذه الشركة يعطي 250002."""
    code = normalize_company_code(company_code)
    if not code:
        raise ValueError("رمز الشركة غير صالح؛ يجب أن يكون رقماً بين 01 و 99.")
    max_seq = 0
    for inv in load_invoices():
        invno = str(inv.get("invno", "")).strip()
        if len(invno) == 6 and invno.isdigit() and invno[:2] == code:
            max_seq = max(max_seq, int(invno[2:]))
    next_seq = max_seq + 1
    if next_seq > 9999:
        raise ValueError("تم الوصول للحد الأقصى لعدد فواتير هذه الشركة (9999).")
    return f"{code}{next_seq:04d}"


def next_receipt_number(company_code):
    """يولّد رقم سند القبض التالي لعميل معين بنفس أسلوب ترقيم الفواتير:
    رمز الشركة (رقمان) + رقم تسلسلي خاص بسندات قبض هذا العميل (4 أرقام).
    مثال: الرمز 01 مع سند القبض الثالث لهذا العميل يعطي 010003."""
    code = normalize_company_code(company_code)
    if not code:
        raise ValueError("رمز الشركة غير صالح؛ يجب أن يكون رقماً بين 01 و 99.")
    max_seq = 0
    for p in load_payments():
        rid = str(p.get("id", "")).strip()
        if len(rid) == 6 and rid.isdigit() and rid[:2] == code:
            max_seq = max(max_seq, int(rid[2:]))
    next_seq = max_seq + 1
    if next_seq > 9999:
        raise ValueError("تم الوصول للحد الأقصى لعدد سندات القبض لهذه الشركة (9999).")
    return f"{code}{next_seq:04d}"


def save_invoice(d, editing_invno=None):
    """يضيف فاتورة جديدة أو يحدّث فاتورة موجودة، ويحفظ الملف مباشرة على القرص."""
    ensure_workbook()
    wb = load_workbook(DATA_FILE)
    ws = wb[SHEET_INVOICES]
    target_row = None
    if editing_invno is not None:
        for r in range(2, ws.max_row + 1):
            if str(ws.cell(row=r, column=1).value) == str(editing_invno):
                target_row = r
                break
    row_vals = invoice_dict_to_row(d)
    if target_row is None:
        target_row = ws.max_row + 1
        if ws.cell(row=1, column=1).value is None:
            target_row = 1  # لن يحدث عادة لأن ensure_workbook يكتب العناوين
    for c, val in enumerate(row_vals, start=1):
        ws.cell(row=target_row, column=c, value=val)
    _safe_save(wb)


def delete_invoice(invno):
    ensure_workbook()
    wb = load_workbook(DATA_FILE)
    ws = wb[SHEET_INVOICES]
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value) == str(invno):
            ws.delete_rows(r, 1)
            break
    _safe_save(wb)


def load_customers():
    ensure_workbook()
    wb = load_workbook(DATA_FILE)
    ws = wb[SHEET_CUSTOMERS]
    rows = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=1).value
        if name in (None, ""):
            continue
        rows.append({
            "name": name,
            "phone": ws.cell(row=r, column=2).value or "",
            "address": ws.cell(row=r, column=3).value or "",
            "notes": ws.cell(row=r, column=4).value or "",
            "company_code": normalize_company_code(ws.cell(row=r, column=5).value),
        })
    return rows


def save_customer(c, editing_name=None):
    ensure_workbook()
    wb = load_workbook(DATA_FILE)
    ws = wb[SHEET_CUSTOMERS]
    target_row = None
    if editing_name is not None:
        for r in range(2, ws.max_row + 1):
            if ws.cell(row=r, column=1).value == editing_name:
                target_row = r
                break
    if target_row is None:
        target_row = ws.max_row + 1
    ws.cell(row=target_row, column=1, value=c["name"])
    ws.cell(row=target_row, column=2, value=c["phone"])
    ws.cell(row=target_row, column=3, value=c["address"])
    ws.cell(row=target_row, column=4, value=c["notes"])
    ws.cell(row=target_row, column=5, value=c.get("company_code", ""))
    _safe_save(wb)


def delete_customer(name):
    ensure_workbook()
    wb = load_workbook(DATA_FILE)
    ws = wb[SHEET_CUSTOMERS]
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value == name:
            ws.delete_rows(r, 1)
            break
    _safe_save(wb)


# ------------------------------------------------------------------
# الدفعات (المبالغ المدفوعة من العملاء) وحساب الأرصدة المستحقة
# ------------------------------------------------------------------
def load_payments():
    """يقرأ جميع الدفعات المسجّلة ويعيدها كقائمة قواميس."""
    ensure_workbook()
    wb = load_workbook(DATA_FILE)
    ws = wb[SHEET_PAYMENTS]
    rows = []
    for r in range(2, ws.max_row + 1):
        pid = ws.cell(row=r, column=1).value
        if pid in (None, ""):
            continue
        applied_raw = ws.cell(row=r, column=6).value or ""
        applied = [x.strip() for x in str(applied_raw).split(",") if x.strip()]
        rows.append({
            "id": pid,
            "customer": ws.cell(row=r, column=2).value or "",
            "date": ws.cell(row=r, column=3).value or "",
            "amount": ws.cell(row=r, column=4).value or 0,
            "notes": ws.cell(row=r, column=5).value or "",
            "applied_invoices": applied,
        })
    return rows


def save_payment(p, editing_id=None):
    """يضيف دفعة جديدة أو يحدّث دفعة موجودة (برقم السند)، ويحفظ الملف مباشرة على القرص.
    عند إضافة دفعة جديدة، يُولَّد لها تلقائياً رقم سند قبض بصيغة رمز الشركة (رقمان) +
    رقم تسلسلي (4 أرقام) خاص بهذا العميل، مطابقاً لأسلوب ترقيم الفواتير.
    يعيد رقم السند الخاص بالدفعة المحفوظة."""
    ensure_workbook()
    wb = load_workbook(DATA_FILE)
    ws = wb[SHEET_PAYMENTS]
    target_row = None
    if editing_id is not None:
        for r in range(2, ws.max_row + 1):
            if str(ws.cell(row=r, column=1).value) == str(editing_id):
                target_row = r
                break
    customer_name = p.get("customer", "")
    if target_row is None:
        target_row = ws.max_row + 1
        if ws.cell(row=1, column=1).value is None:
            target_row = 1
        company_code = ""
        for c in load_customers():
            if c["name"] == customer_name:
                company_code = c.get("company_code", "")
                break
        if not company_code:
            raise ValueError(f'العميل "{customer_name}" ليس له رمز شركة؛ لا يمكن توليد رقم سند القبض تلقائياً.')
        pid = next_receipt_number(company_code)
    else:
        original_customer = ws.cell(row=target_row, column=2).value or ""
        if original_customer != customer_name:
            company_code = ""
            for c in load_customers():
                if c["name"] == customer_name:
                    company_code = c.get("company_code", "")
                    break
            if not company_code:
                raise ValueError(f'العميل "{customer_name}" ليس له رمز شركة؛ لا يمكن توليد رقم سند القبض تلقائياً.')
            pid = next_receipt_number(company_code)
        else:
            pid = editing_id
    ws.cell(row=target_row, column=1, value=pid)
    ws.cell(row=target_row, column=2, value=p.get("customer", ""))
    ws.cell(row=target_row, column=3, value=p.get("date", ""))
    ws.cell(row=target_row, column=4, value=round(float(p.get("amount") or 0), 3))
    ws.cell(row=target_row, column=5, value=p.get("notes", ""))
    applied = p.get("applied_invoices") or []
    ws.cell(row=target_row, column=6, value=", ".join(str(x).strip() for x in applied if str(x).strip()))
    _safe_save(wb)
    return pid


def delete_payment(payment_id):
    ensure_workbook()
    wb = load_workbook(DATA_FILE)
    ws = wb[SHEET_PAYMENTS]
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value) == str(payment_id):
            ws.delete_rows(r, 1)
            break
    _safe_save(wb)


def customer_balances(invoices=None, payments=None):
    """يحسب لكل عميل: إجمالي الفواتير، إجمالي المدفوع، والرصيد المستحق (فواتير - مدفوع).
    يعيد قائمة قواميس مرتبة أبجدياً حسب اسم العميل: {customer, invoiced, paid, balance}."""
    invoices = invoices if invoices is not None else load_invoices()
    payments = payments if payments is not None else load_payments()

    invoiced_by_customer = {}
    for inv in invoices:
        name = inv.get("customer") or ""
        if not name:
            continue
        invoiced_by_customer[name] = invoiced_by_customer.get(name, 0.0) + float(inv.get("total") or 0)

    paid_by_customer = {}
    for p in payments:
        name = p.get("customer") or ""
        if not name:
            continue
        paid_by_customer[name] = paid_by_customer.get(name, 0.0) + float(p.get("amount") or 0)

    all_names = sorted(set(invoiced_by_customer) | set(paid_by_customer))
    result = []
    for name in all_names:
        invoiced = round(invoiced_by_customer.get(name, 0.0), 3)
        paid = round(paid_by_customer.get(name, 0.0), 3)
        result.append({
            "customer": name,
            "invoiced": invoiced,
            "paid": paid,
            "balance": round(invoiced - paid, 3),
        })
    return result


def allocate_payment_to_invoices(inv_by_no, amount, applied_invoice_numbers):
    """يوزّع مبلغ دفعة على فواتير محددة بالترتيب المعطى، ويعيد المبلغ المتبقي.
    كل فاتورة تُخصم منها ما يلزم من المبلغ حتى ينفد المبلغ أو تُسدَّد الفاتورة بالكامل."""
    remaining = round(float(amount or 0), 3)
    allocations = []
    for no in applied_invoice_numbers:
        if remaining <= 0.0005:
            break
        inv = inv_by_no.get(no)
        if not inv:
            continue
        due = round(float(inv.get("total", 0) or 0) - float(inv.get("paid", 0) or 0), 3)
        if due <= 0.0005:
            continue
        pay_amt = round(min(due, remaining), 3)
        inv["paid"] = round(float(inv.get("paid", 0) or 0) + pay_amt, 3)
        remaining = round(remaining - pay_amt, 3)
        allocations.append((no, pay_amt))
    return remaining, allocations


def invoice_payment_status(invoices=None, payments=None):
    """يحسب لكل فاتورة: الإجمالي، المدفوع منها، المتبقي، والحالة (مسددة/جزئية/غير مسددة).

    آلية التوزيع:
      - إن حدّدت الدفعة فواتير معيّنة ("الفواتير المسددة")، يُطبَّق مبلغها عليها بالترتيب
        الذي اختارها به المستخدم حتى ينفد المبلغ أو تُسدَّد الفواتير المحددة بالكامل.
        أي مبلغ متبقٍ بعد ذلك يُضاف إلى "الرصيد العام" لهذا العميل.
      - الدفعات التي لا تحدد فواتير (دفعة إجمالية من إجمالي المستحق) تُضاف مباشرة
        إلى "الرصيد العام" لهذا العميل.
      - يُطبَّق "الرصيد العام" لكل عميل تلقائياً على أقدم فواتيره غير المسددة أولاً (FIFO)
        حسب التاريخ، وهذا يحاكي الفواتير التي تُسدَّد كجزء من دفعة مجملة غير مخصصة.

    يعيد قاموساً: {رقم_الفاتورة: {"total", "paid", "balance", "status"}}.
    """
    invoices = invoices if invoices is not None else load_invoices()
    payments = payments if payments is not None else load_payments()

    inv_by_no = {}
    by_customer = {}
    for inv in invoices:
        no = str(inv.get("invno", ""))
        if not no:
            continue
        inv_by_no[no] = {"total": float(inv.get("total") or 0), "paid": 0.0,
                          "customer": inv.get("customer", ""), "date": str(inv.get("date", ""))}
        by_customer.setdefault(inv.get("customer", ""), []).append(no)

    for name in by_customer:
        by_customer[name].sort(key=lambda no: (inv_by_no[no]["date"], no))

    general_pool = {}
    for p in payments:
        amount = float(p.get("amount") or 0)
        customer = p.get("customer", "")
        applied = [str(a).strip() for a in (p.get("applied_invoices") or []) if str(a).strip()]
        if applied:
            remaining, _ = allocate_payment_to_invoices(inv_by_no, amount, applied)
            if remaining > 0.0005:
                general_pool[customer] = general_pool.get(customer, 0.0) + remaining
        else:
            general_pool[customer] = general_pool.get(customer, 0.0) + amount

    for customer, pool in general_pool.items():
        remaining = pool
        for no in by_customer.get(customer, []):
            if remaining <= 0.0005:
                break
            inv = inv_by_no[no]
            due = inv["total"] - inv["paid"]
            if due <= 0.0005:
                continue
            pay_amt = min(due, remaining)
            inv["paid"] += pay_amt
            remaining -= pay_amt

    result = {}
    for no, inv in inv_by_no.items():
        total = round(inv["total"], 3)
        paid = round(inv["paid"], 3)
        balance = round(total - paid, 3)
        if balance <= 0.0005:
            status = "paid"
        elif paid > 0.0005:
            status = "partial"
        else:
            status = "unpaid"
        result[no] = {"total": total, "paid": paid, "balance": balance, "status": status}
    return result


def get_dashboard_stats(invoices=None, payments=None, customers=None):
    """يحسب مؤشرات الأداء الرئيسية للوحة المعلومات: عدد الفواتير، عدد العملاء،
    إجمالي المبالغ المستحقة، إجمالي المدفوع، الرصيد المستحق، حالة كل فاتورة، وأكبر المدينين."""
    invoices = invoices if invoices is not None else load_invoices()
    payments = payments if payments is not None else load_payments()
    customers = customers if customers is not None else load_customers()

    total_invoiced = round(sum(float(i.get("total") or 0) for i in invoices), 3)
    total_paid = round(sum(float(p.get("amount") or 0) for p in payments), 3)
    total_outstanding = round(total_invoiced - total_paid, 3)

    status_map = invoice_payment_status(invoices, payments)
    paid_count = sum(1 for v in status_map.values() if v["status"] == "paid")
    partial_count = sum(1 for v in status_map.values() if v["status"] == "partial")
    unpaid_count = sum(1 for v in status_map.values() if v["status"] == "unpaid")

    balances = customer_balances(invoices, payments)
    top_debtors = sorted([b for b in balances if b["balance"] > 0.0005],
                          key=lambda b: -b["balance"])[:5]

    return {
        "invoice_count": len(invoices),
        "customer_count": len(customers),
        "payment_count": len(payments),
        "total_invoiced": total_invoiced,
        "total_paid": total_paid,
        "total_outstanding": total_outstanding,
        "paid_count": paid_count,
        "partial_count": partial_count,
        "unpaid_count": unpaid_count,
        "top_debtors": top_debtors,
        "avg_invoice": round(total_invoiced / len(invoices), 3) if invoices else 0.0,
    }


# ------------------------------------------------------------------
# واجهة المستخدم
# ------------------------------------------------------------------
