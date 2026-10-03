# -*- coding: utf-8 -*-
"""
core.data.storage — عمليات القراءة/الكتابة العامة لكل الجداول.

يوزّع كل عملية على المحرك المختار في runtime (SQLite افتراضياً أو Excel
للتوافق) ، ويضيف طبقة الحماية من فقدان البيانات: حفظ ذرّي، نسخ احتياطية يومية،
سجل حركات، وقفل نسخة واحدة.

لا يعرف شيثاً عن قعاد العمل — تلك في core.data.repos.
"""

import csv
import os
import shutil
from datetime import date, datetime

from openpyxl import Workbook, load_workbook

from core.data import runtime, sqlite
from core.data.config import (
    AUDIT_DIR_NAME, AUDIT_HEADERS, BACKUP_DIR_NAME, HEADER_ALIGNMENT, HEADER_FILL,
    HEADER_FONT, MAX_BACKUPS, SHEETS, headers_for,
)
from core.data.validation import safe_int


def _write_header(ws, fields):
    for i, (_, label) in enumerate(fields, start=1):
        c = ws.cell(row=1, column=i, value=label)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = HEADER_ALIGNMENT
    ws.freeze_panes = "A2"


def ensure_workbook():
    """
    يجهّز مخزن البيانات:
      - SQLite (الافتراضي): ينشئ قاعدة البيانات وكل جداولها وأعمدتها الناقصة.
        عند أول إقلاع والقاعدة فارغة وكان ملف Excel موجوداً بجانبها، تُستورد
        بياناته تلقائياً (بعد نسخة احتياطية منه) — الترحيل دون أي خطوة يدوية.
      - Excel (وضع التوافق القديم RESTAURANT_BACKEND=excel): يُنشئ الملف وكل
        أوراقه إذا لم يكن موجوداً، ويضيف أي ورقة ناقصة إلى ملف موجود مسبقاً.
    """
    if runtime.is_sqlite():
        sqlite.ensure()
        _auto_import_excel_once()
        return
    if os.path.exists(runtime.DATA_FILE):
        if not runtime._schema_checked:
            _add_missing_sheets()
            runtime._schema_checked = True
        return
    wb = Workbook()
    first = True
    for key, sheet_name in SHEETS.items():
        ws = wb.active if first else wb.create_sheet(sheet_name)
        if first:
            ws.title = sheet_name
            first = False
        _write_header(ws, headers_for(key))
    save_workbook(wb)


def _add_missing_sheets():
    """يضيف الجداول الناقصة (بعناوين أعمدتها) إلى ملف بيانات موجود."""
    try:
        wb = load_workbook(runtime.DATA_FILE)
    except Exception:
        return 0
    added = 0
    for key, sheet_name in SHEETS.items():
        if sheet_name not in wb.sheetnames:
            _write_header(wb.create_sheet(sheet_name), headers_for(key))
            added += 1
            log_action("إضافة جدول ناقص", key, details=sheet_name)
    # ترقية أعمدة ملف قديم: أي عمود أُضيف إلى مخطط الجداول في إصدار أحدث
    # تُكتب ترويسته في نهاية الورقة دون تحريك الأعمدة القائمة (لا يُفقد أي بيان).
    added += _add_missing_columns(wb)
    if added:
        save_workbook(wb)
    return added


def _add_missing_columns(wb):
    """
    يكتب ترويسات الأعمدة التي أُضيفت إلى المخطط في إصدارات أحدث ولا تزال
    فارغة في ملف قديم. يفترض أن الإضافات الجديدة تُلحق دائماً في نهاية قائمة
    الحقول حتى لا تنزاح مواضع الأعمدة الموجودة.
    """
    added = 0
    for key, sheet_name in SHEETS.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        for i, (_, label) in enumerate(headers_for(key), start=1):
            if ws.cell(row=1, column=i).value in (None, ""):
                c = ws.cell(row=1, column=i, value=label)
                c.font = HEADER_FONT
                c.fill = HEADER_FILL
                c.alignment = HEADER_ALIGNMENT
                added += 1
                log_action("ترقية مخطط", key, details=label)
    return added


def next_id(rows):
    ids = [safe_int(r.get("id")) for r in rows if str(r.get("id", "")).strip() != ""]
    return (max(ids) if ids else 0) + 1


def save_workbook(wb):
    """
    حفظ آمن لملف البيانات: يُكتب في ملف مؤقت ثم يُستبدل الملف الأصلي، مع
    الاحتفاظ بنسخة (runtime.BAK_FILE) من الإصدار السابق مباشرة.
    بهذه الطريقة لا يمكن أن ينكسر ملف البيانات إذا توقف البرنامج أثناء الحفظ.
    (محرك SQLite لا يحتاجها: كل عملية معاملة مستقلة.)
    """
    if runtime.is_sqlite():
        return
    tmp_path = runtime.DATA_FILE + ".tmp"
    wb.save(tmp_path)
    if os.path.exists(runtime.DATA_FILE):
        try:
            shutil.copy2(runtime.DATA_FILE, runtime.BAK_FILE)
        except OSError:
            pass
    os.replace(tmp_path, runtime.DATA_FILE)


def log_action(action, sheet_key="", key="", details=""):
    """
    يوثّق كل عملية إضافة/تعديل/حذف في ملف CSV شهري داخل logs/.
    ملف السجل منفصل عن ملف Excel حتى لا يتغير شكل بيانات المستخدم، وأي فشل
    في التسجيل لا يجوز أن يوقف عملية الحفظ.
    """
    try:
        path = os.path.join(runtime.sub_dir(AUDIT_DIR_NAME),
                            "audit_%s.csv" % date.today().strftime("%Y-%m"))
        new_file = not os.path.exists(path)
        with open(path, "a", encoding="utf-8-sig", newline="") as fh:
            writer = csv.writer(fh)
            if new_file:
                writer.writerow(AUDIT_HEADERS)
            writer.writerow([datetime.now().isoformat(timespec="seconds"), action,
                             sheet_key, key, details])
    except OSError:
        pass


def _prune_backups(keep=MAX_BACKUPS):
    """يحذف أقدم النسخ الاحتياطية مع الإبقاء على آخر (keep) نسخة."""
    backup_dir = runtime.sub_dir(BACKUP_DIR_NAME, create=False)
    if not os.path.isdir(backup_dir):
        return 0
    files = sorted(name for name in os.listdir(backup_dir)
                   if name.startswith("restaurant_data_") and name.endswith(".xlsx"))
    removed = 0
    for name in files[:-keep] if keep > 0 else files:
        try:
            os.remove(os.path.join(backup_dir, name))
            removed += 1
        except OSError:
            pass
    return removed


def backup_data_file(force=False):
    """
    نسخة احتياطية يومية من مخزن البيانات في backups/.
    تعيد مسار النسخة، أو None إذا كانت نسخة اليوم موجودة بالفعل (force=True يتجاوز ذلك).
    """
    if runtime.is_sqlite():
        return sqlite.backup_data_file(force)
    if not os.path.exists(runtime.DATA_FILE):
        return None
    target = os.path.join(runtime.sub_dir(BACKUP_DIR_NAME),
                          "restaurant_data_%s.xlsx" % date.today().strftime("%Y-%m-%d"))
    if os.path.exists(target) and not force:
        return None
    shutil.copy2(runtime.DATA_FILE, target)
    _prune_backups()
    log_action("نسخة احتياطية", details=os.path.basename(target))
    return target


def startup_maintenance():
    """
    يُستدعى عند تشغيل البرنامج: تجهيز ملف البيانات ثم أخذ نسخة احتياطية يومية.
    تعيد مسار النسخة الاحتياطية إن أُخذت، أو None.
    """
    ensure_workbook()
    return backup_data_file()


def acquire_instance_lock():
    """
    قفل يمنع تشغيل نسختين من البرنامج على نفس ملف البيانات معاً، لأن تشغيل
    نسختين هو السبب الأول لفقدان البيانات في ملفات Excel (آخر حفظ يلغي الآخر).
    تعيد True إذا أمكن المتابعة، و False إذا كان البرنامج مفتوحاً في مكان آخر.
    """
    if runtime._lock_handle is not None:
        return True
    lock_path = (sqlite.DB_FILE if runtime.is_sqlite() else runtime.DATA_FILE) + ".lock"
    try:
        handle = open(lock_path, "a+")
    except OSError:
        return True                      # تعذّر إنشاء ملف القفل: لا نمنع التشغيل
    try:
        if os.name == "nt":
            import msvcrt
            handle.seek(0)
            msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
        else:
            import fcntl
            fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
    except ImportError:
        pass                             # لا تتوفر أدوات القفل: نتابع بدون منع
    except OSError:
        handle.close()
        return False
    try:
        handle.seek(0)
        handle.truncate()
        handle.write("pid=%s" % os.getpid())
        handle.flush()
    except OSError:
        pass
    runtime._lock_handle = handle
    return True


def release_instance_lock():
    """يفتح القفل عند إغلاق البرنامج (النظام يحرره تلقائياً عند انتهاء العملية)."""
    if runtime._lock_handle is None:
        return
    try:
        if os.name == "nt":
            import msvcrt
            runtime._lock_handle.seek(0)
            msvcrt.locking(runtime._lock_handle.fileno(), msvcrt.LK_UNLCK, 1)
        else:
            import fcntl
            fcntl.flock(runtime._lock_handle.fileno(), fcntl.LOCK_UN)
    except (ImportError, OSError):
        pass
    try:
        runtime._lock_handle.close()
    except OSError:
        pass
    runtime._lock_handle = None


def _load_rows(sheet_key, key_field):
    if runtime.is_sqlite():
        return sqlite.load_rows(sheet_key)
    ensure_workbook()
    wb = load_workbook(runtime.DATA_FILE)
    ws = wb[SHEETS[sheet_key]]
    fields = headers_for(sheet_key)
    rows = []
    for r in range(2, ws.max_row + 1):
        first_val = ws.cell(row=r, column=1).value
        if first_val in (None, ""):
            continue
        row_dict = {}
        for i, (key, _) in enumerate(fields, start=1):
            v = ws.cell(row=r, column=i).value
            row_dict[key] = v if v is not None else ""
        rows.append(row_dict)
    return rows


def _save_row(sheet_key, key_field, data, editing_key=None):
    if runtime.is_sqlite():
        sqlite.save_row(sheet_key, data, key_field, editing_key)
        log_action("تعديل" if editing_key is not None else "إضافة", sheet_key,
                   data.get(key_field, ""))
        return
    ensure_workbook()
    wb = load_workbook(runtime.DATA_FILE)
    ws = wb[SHEETS[sheet_key]]
    fields = headers_for(sheet_key)
    target_row = None
    search_val = editing_key if editing_key is not None else data.get(key_field)
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value) == str(search_val):
            target_row = r
            break
    if target_row is None:
        target_row = ws.max_row + 1
    for i, (key, _) in enumerate(fields, start=1):
        ws.cell(row=target_row, column=i, value=data.get(key, ""))
    save_workbook(wb)
    log_action("تعديل" if editing_key is not None else "إضافة", sheet_key,
               data.get(key_field, ""))


def _save_row_by_name(sheet_key, data, editing_name):
    """حفظ صف يستخدم فيه أول عمود (الاسم) كمفتاح نصي بدل رقم تسلسلي."""
    if runtime.is_sqlite():
        key_field = headers_for(sheet_key)[0][0]
        sqlite.save_row(sheet_key, data, key_field, editing_name)
        log_action("تعديل" if editing_name is not None else "إضافة", sheet_key,
                   data.get(key_field, ""))
        return
    ensure_workbook()
    wb = load_workbook(runtime.DATA_FILE)
    ws = wb[SHEETS[sheet_key]]
    fields = headers_for(sheet_key)
    target_row = None
    search_val = editing_name if editing_name is not None else data.get(fields[0][0])
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value) == str(search_val):
            target_row = r
            break
    if target_row is None:
        target_row = ws.max_row + 1
    for i, (key, _) in enumerate(fields, start=1):
        ws.cell(row=target_row, column=i, value=data.get(key, ""))
    save_workbook(wb)
    log_action("تعديل" if editing_name is not None else "إضافة", sheet_key,
               data.get(fields[0][0], ""))


def _delete_row(sheet_key, key_field, key_value):
    if runtime.is_sqlite():
        if sqlite.delete_row(sheet_key, key_value):
            log_action("حذف", sheet_key, key_value)
        return
    ensure_workbook()
    wb = load_workbook(runtime.DATA_FILE)
    ws = wb[SHEETS[sheet_key]]
    deleted = False
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value) == str(key_value):
            ws.delete_rows(r, 1)
            deleted = True
            break
    if deleted:
        save_workbook(wb)
        log_action("حذف", sheet_key, key_value)


# اسماء عامة يُستخدمها المستودعات (names without the underscore prefix)
load_rows = _load_rows
save_row = _save_row
save_row_by_name = _save_row_by_name
delete_row = _delete_row


# ------------------------------------------------------------------
# أدوات التبادل مع Excel (يُصبح المحرك مرتبًط بملف معلومات)
# ------------------------------------------------------------------
def _auto_import_excel_once():
    """
    الاستيراد التلقائي لمرة واحدة: عند أول إقلاع على SQLite والقاعدة فارغة
    وكان ملف Excel القديم (runtime.DATA_FILE) موجوداً، تُؤخذ نسخة احتياطية منه ثم
    تُستورد بياناته — يعمل مرة واحدة فقط (أي كتابة لاحقة تجعل القاعدة
    غير فارغة فلا يتكرر). يعيد عدد الصفوف المستوردة أو 0.
    """
    try:
        if not os.path.exists(runtime.DATA_FILE):
            return 0
        if not sqlite.is_empty():
            return 0
        backup_dir = os.path.join(os.path.dirname(os.path.abspath(runtime.DATA_FILE)),
                                  BACKUP_DIR_NAME)
        os.makedirs(backup_dir, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        shutil.copy2(runtime.DATA_FILE, os.path.join(
            backup_dir, "restaurant_data_pre_sqlite_%s.xlsx" % stamp))
        counts = sqlite.import_excel(runtime.DATA_FILE)
        total = sum(counts.values())
        log_action("استيراد تلقائي", "", "",
                   "من Excel إلى SQLite: %d صف في %d جدول"
                   % (total, len(counts)))
        return total
    except Exception:
        return 0


def export_excel(path=None):
    """
    تصدير المخزن الحالي (SQLite) إلى ملف Excel بنفس أسماء الأوراق وترتيب
    الأعمدة — نسخة للتبادل أو للدراسة في Excel. في وضع Excel القديم لا حاجة
    للتصدير (البيانات أصلاً في ملف Excel) فيُعاد مسار ملف البيانات نفسه.
    """
    if not runtime.is_sqlite():
        if path is None:
            return runtime.DATA_FILE
        shutil.copy2(runtime.DATA_FILE, path)
        return path
    return sqlite.export_excel(path)


def import_excel(path):
    """
    ترحيل ملف Excel (بنفس مخطط الأوراق) إلى قاعدة SQLite: يستبدل محتوى
    الجداول بمحتوى الملف. تعيد قاموس الجدول -> عدد الصفوف المستوردة.
    (في وضع Excel القديم لا شيء يُستورد — البيانات هناك أصلاً.)
    """
    if not runtime.is_sqlite():
        return {}
    return sqlite.import_excel(path)


# توصيل وصف الجداول ودالفة السجل بمحرك SQLite عند الإقلاع
# (المحرك لا يستورد هذا أبداً—لأن يتلقي وصفهما عند الإقلاع)
if runtime.is_sqlite():
    sqlite.configure(SHEETS, headers_for, log_action)

