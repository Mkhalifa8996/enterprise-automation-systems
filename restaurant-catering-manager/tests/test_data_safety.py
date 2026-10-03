# -*- coding: utf-8 -*-
"""
اختبارات حماية البيانات: الحفظ الآمن، النسخ الاحتياطية، سجل الحركات،
قفل النسخة الواحدة، والإصلاح التلقائي لملف ناقص جدول.

    python tests/test_data_safety.py
"""
import os
import shutil
import subprocess
import sys
import tempfile

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE)

# مجموعة السلامة مخصصة لمحرك Excel القديم (تختبر النسخ والحفظ الآمن).
SANDBOX = os.path.join(tempfile.gettempdir(), "restaurant_safety_tests")
shutil.rmtree(SANDBOX, ignore_errors=True)
os.makedirs(SANDBOX, exist_ok=True)
os.environ["RESTAURANT_BACKEND"] = "excel"
os.environ["RESTAURANT_DATA_FILE"] = os.path.join(SANDBOX, "restaurant_data.xlsx")

import restaurant_data as rd

rd._schema_checked = False

passed, failed = [], []


def case(label):
    def deco(fn):
        try:
            fn()
            passed.append(label)
            print("  [ OK ] %s" % label)
        except Exception as exc:                 # noqa: BLE001
            failed.append((label, exc))
            print("  [FAIL] %s -> %s: %s" % (label, type(exc).__name__, exc))
        return fn
    return deco


def assert_true(cond, label):
    assert cond, label


def assert_eq(actual, expected, label):
    assert actual == expected, "%s: المتوقع %r والفعل %r" % (label, expected, actual)


def _backup_dir():
    return os.path.join(SANDBOX, rd.BACKUP_DIR_NAME)


def _audit_path():
    return os.path.join(SANDBOX, rd.AUDIT_DIR_NAME,
                        "audit_%s.csv" % rd.date.today().strftime("%Y-%m"))


print("1) النسخ الاحتياطية اليومية")


@case("startup_maintenance ينشئ ملف البيانات ونسخة اليوم")
def _t_daily_backup():
    path = rd.startup_maintenance()
    assert_true(os.path.exists(rd.DATA_FILE), "ملف البيانات أُنشئ")
    assert_true(path is not None and os.path.exists(path), "نسخة اليوم أُخذت")
    assert_eq(os.path.basename(path),
              "restaurant_data_%s.xlsx" % rd.date.today().strftime("%Y-%m-%d"),
              "اسم نسخة اليوم")


@case("لا تُكرَّر نسخة نفس اليوم إلا عند الطلب")
def _t_daily_backup_once():
    assert_eq(rd.backup_data_file(), None, "النسخة الثانية في نفس اليوم لا تُنشأ")
    forced = rd.backup_data_file(force=True)
    assert_true(forced is not None and os.path.exists(forced), "force=True يُنشئ النسخة")


@case("تقليم النسخ يحتفظ بأحدث %d نسخة" % rd.MAX_BACKUPS)
def _t_prune():
    for i in range(1, rd.MAX_BACKUPS + 6):
        open(os.path.join(_backup_dir(), "restaurant_data_2025-01-%03d.xlsx" % i), "wb").close()
    rd._prune_backups()
    names = sorted(n for n in os.listdir(_backup_dir())
                   if n.startswith("restaurant_data_") and n.endswith(".xlsx"))
    assert_eq(len(names), rd.MAX_BACKUPS, "عدد النسخ المتبقية")
    assert_true("restaurant_data_2025-01-001.xlsx" not in names, "أقدم نسخة حُذفت")
    assert_true("restaurant_data_2025-01-035.xlsx" in names, "أحدث نسخة محفوظة")


print("2) الحفظ الآمن (atomic save)")


@case("الحفظ لا يترك ملفاً مؤقتاً وينشئ .bak من الإصدار السابق")
def _t_atomic_save():
    rd.save_customer({"name": "عميل الحفظ الآمن", "type": "شركة"})
    assert_true(not os.path.exists(rd.DATA_FILE + ".tmp"), "لا يبقى ملف .tmp")
    bak = rd.BAK_FILE
    assert_true(os.path.exists(bak), "نسخة .bak أُنشئت")
    from openpyxl import load_workbook
    ws = load_workbook(bak)[rd.SHEETS["customers"]]
    names = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert_true("عميل الحفظ الآمن" not in names, "الـ.bak يخص الإصدار السابق للحفظ")
    assert_true(any(c["name"] == "عميل الحفظ الآمن" for c in rd.load_customers()),
                "العميل الجديد في الملف الحالي")


@case("ملف .bak صالح ويُفتح كملف Excel")
def _t_bak_readable():
    from openpyxl import load_workbook
    wb = load_workbook(rd.BAK_FILE)
    assert_true(rd.SHEETS["customers"] in wb.sheetnames, "الـ.bak ملف Excel سليم")


print("3) سجل الحركات (audit log)")


@case("السجل يوثّق الإضافة والتعديل والحذف")
def _t_audit():
    rd.save_customer({"name": "عميل السجل", "type": "شركة"})
    row = [c for c in rd.load_customers() if c["name"] == "عميل السجل"][0]
    row["phone"] = "999"
    rd.save_customer(row, editing_name="عميل السجل")
    rd.delete_customer("عميل السجل")
    assert_true(os.path.exists(_audit_path()), "ملف السجل أُنشئ")
    text = open(_audit_path(), encoding="utf-8-sig").read()
    for header in rd.AUDIT_HEADERS:
        assert_true(header in text, "عنوان السجل: %s" % header)
    for action in ("إضافة", "تعديل", "حذف"):
        assert_true(action in text, "الإجراء %s مُسجّل" % action)


@case("السجل يوثّق إصدار الفاتورة")
def _t_audit_invoice():
    rd.save_institutional_order({"date": rd.date.today().isoformat(), "customer": "شركة الخليج",
                                 "meal_type": "غداء", "quantity": "10", "unit_price": "1.000"})
    order = rd.load_institutional_orders()[0]
    rd.create_invoice_for_customer("شركة الخليج", [order["id"]], [])
    text = open(_audit_path(), encoding="utf-8-sig").read()
    assert_true("إصدار فاتورة" in text, "إصدار الفاتورة مُسجّل")